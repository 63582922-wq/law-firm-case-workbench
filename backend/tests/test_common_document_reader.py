from __future__ import annotations

from hashlib import sha256
import io
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from uuid import uuid4
import zipfile

from case_kernel.common_document_reader import (
    CommonDocumentFormat,
    CommonDocumentReadingBlocked,
    DocumentCandidateKind,
    DocumentReadBudget,
    DocumentReviewStatus,
    MaterializedDocumentSource,
    read_materialized_common_document,
)


class CommonDocumentReaderTests(unittest.TestCase):
    def source(
        self,
        root: Path,
        name: str,
        content: bytes,
        document_format: CommonDocumentFormat,
    ) -> MaterializedDocumentSource:
        path = root / name
        path.write_bytes(content)
        path.chmod(0o600)
        return MaterializedDocumentSource(
            source_object_id=str(uuid4()),
            source_object_version="version-opaque-1",
            materialization_root=root,
            path=path,
            byte_size=len(content),
            content_sha256=sha256(content).hexdigest(),
            admitted_format=document_format,
        )

    @staticmethod
    def package(entries: dict[str, str | bytes], *, compression: int = zipfile.ZIP_STORED) -> bytes:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=compression) as archive:
            for name, value in entries.items():
                archive.writestr(name, value)
        return buffer.getvalue()

    def test_docx_returns_source_bound_paragraph_table_hidden_and_tracked_candidates(self) -> None:
        content = self.package(
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": """
                    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                      <w:body>
                        <w:p><w:r><w:t>答辩意见</w:t></w:r></w:p>
                        <w:p><w:ins><w:r><w:t>新增事实</w:t></w:r></w:ins></w:p>
                        <w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>隐藏备注</w:t></w:r></w:p>
                        <w:tbl><w:tr><w:tc><w:p><w:r><w:t>金额100元</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
                      </w:body>
                    </w:document>
                """,
            }
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "case.docx", content, CommonDocumentFormat.DOCX)
            source_hash = source.content_sha256
            result = read_materialized_common_document(source)
        self.assertEqual(result.source_sha256, source_hash)
        self.assertEqual(result.review_status, DocumentReviewStatus.NEEDS_REVIEW)
        self.assertTrue(all(item.review_status is DocumentReviewStatus.NEEDS_REVIEW for item in result.candidates))
        texts = {item.text: item for item in result.candidates}
        self.assertIn("答辩意见", texts)
        self.assertIn("金额100元", texts)
        self.assertEqual(texts["金额100元"].kind, DocumentCandidateKind.TABLE_CELL)
        self.assertIn("TRACKED_CHANGE_CONTENT", texts["新增事实"].risk_flags)
        self.assertIn("HIDDEN_CONTENT", texts["隐藏备注"].risk_flags)
        self.assertEqual(texts["金额100元"].location.row, 1)
        self.assertEqual(texts["金额100元"].location.column, 1)
        self.assertNotIn("答辩意见", repr(result))
        self.assertNotIn("金额100元", repr(texts["金额100元"]))

    def test_docx_macro_embed_external_relationship_and_path_traversal_are_blocked(self) -> None:
        cases = (
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>",
                "word/vbaProject.bin": b"macro",
            },
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>",
                "word/embeddings/oleObject1.bin": b"ole",
            },
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>",
                "word/_rels/document.xml.rels": """
                  <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                    <Relationship Id="r1" Type="x" TargetMode="External" Target="https://evil.example/x"/>
                  </Relationships>""",
            },
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>",
                "../outside.txt": "escape",
            },
        )
        for index, entries in enumerate(cases):
            with self.subTest(index=index), TemporaryDirectory() as temporary:
                content = self.package(entries)
                source = self.source(Path(temporary), "bad.docx", content, CommonDocumentFormat.DOCX)
                with self.assertRaises(CommonDocumentReadingBlocked):
                    read_materialized_common_document(source)

    def test_docx_standard_package_relationships_remain_readable(self) -> None:
        try:
            from docx import Document
        except ImportError:  # project dependency is exercised in the full env
            self.skipTest("python-docx dependency is unavailable")
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "real.docx"
            document = Document()
            document.add_paragraph("标准文书")
            document.save(path)
            raw = path.read_bytes()
            source = MaterializedDocumentSource(
                str(uuid4()), "opaque", root, path, len(raw), sha256(raw).hexdigest(),
                CommonDocumentFormat.DOCX,
            )
            result = read_materialized_common_document(source)
        self.assertEqual(result.candidates[0].text, "标准文书")

    def test_xlsx_preserves_formula_literal_and_hidden_content_without_execution(self) -> None:
        content = self.package(
            {
                "[Content_Types].xml": "<Types/>",
                "xl/workbook.xml": """
                  <workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                    <sheets><sheet name="流水" sheetId="1" state="hidden" r:id="rId1"/></sheets>
                  </workbook>""",
                "xl/_rels/workbook.xml.rels": """
                  <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                    <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
                  </Relationships>""",
                "xl/worksheets/sheet1.xml": """
                  <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
                    <cols><col min="2" max="2" hidden="1"/></cols>
                    <sheetData><row r="1" hidden="1"><c r="A1" t="inlineStr"><is><t>已还款</t></is></c>
                    <c r="B1"><f>SUM(B2:B3)</f><v>300</v></c>
                    <c r="C1" t="inlineStr"><is><t>=2+2</t></is></c></row></sheetData>
                  </worksheet>""",
            }
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "ledger.xlsx", content, CommonDocumentFormat.XLSX)
            result = read_materialized_common_document(source)
        cells = {item.location.coordinate: item for item in result.candidates}
        self.assertEqual(cells["B1"].text, "300")
        self.assertIn(("formula_literal", "SUM(B2:B3)"), cells["B1"].attributes)
        self.assertIn("FORMULA_PRESENT_NOT_EVALUATED", cells["B1"].risk_flags)
        self.assertIn("HIDDEN_CONTENT", cells["A1"].risk_flags)
        self.assertIn("FORMULA_LIKE_LITERAL", cells["C1"].risk_flags)
        self.assertIn("HIDDEN_SHEET_PRESENT", result.document_risk_flags)

    def test_xlsx_external_formula_is_blocked(self) -> None:
        content = self.package(
            {
                "[Content_Types].xml": "<Types/>",
                "xl/workbook.xml": """
                  <workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                    <sheets><sheet name="x" sheetId="1" r:id="rId1"/></sheets></workbook>""",
                "xl/_rels/workbook.xml.rels": """
                  <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                    <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
                  </Relationships>""",
                "xl/worksheets/sheet1.xml": """
                  <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
                    <sheetData><row r="1"><c r="A1"><f>WEBSERVICE("https://evil.example")</f><v>0</v></c></row></sheetData>
                  </worksheet>""",
            }
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "bad.xlsx", content, CommonDocumentFormat.XLSX)
            with self.assertRaisesRegex(CommonDocumentReadingBlocked, "formula"):
                read_materialized_common_document(source)

    def test_xlsx_standard_openpyxl_package_remains_readable(self) -> None:
        try:
            from openpyxl import Workbook
        except ImportError:  # project dependency is exercised in the full env
            self.skipTest("openpyxl dependency is unavailable")
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "real.xlsx"
            workbook = Workbook()
            workbook.active.title = "流水"
            workbook.active["A1"] = "已还款"
            workbook.active["B1"] = "=SUM(1,2)"
            workbook.save(path)
            raw = path.read_bytes()
            source = MaterializedDocumentSource(
                str(uuid4()), "opaque", root, path, len(raw), sha256(raw).hexdigest(),
                CommonDocumentFormat.XLSX,
            )
            result = read_materialized_common_document(source)
        cells = {item.location.coordinate: item for item in result.candidates}
        self.assertEqual(cells["A1"].text, "已还款")
        self.assertIn("FORMULA_PRESENT_NOT_EVALUATED", cells["B1"].risk_flags)

    def test_pptx_returns_slide_and_speaker_note_locations(self) -> None:
        content = self.package(
            {
                "[Content_Types].xml": "<Types/>",
                "ppt/presentation.xml": """
                  <p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                   xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                    <p:sldIdLst><p:sldId id="256" show="0" r:id="rId1"/></p:sldIdLst>
                  </p:presentation>""",
                "ppt/_rels/presentation.xml.rels": """
                  <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                    <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
                  </Relationships>""",
                "ppt/slides/slide1.xml": """
                  <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                    <p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>庭审要点</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>
                  </p:sld>""",
                "ppt/notesSlides/notesSlide1.xml": """
                  <p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>律师备注</a:t></p:notes>""",
            }
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "hearing.pptx", content, CommonDocumentFormat.PPTX)
            result = read_materialized_common_document(source)
        by_text = {item.text: item for item in result.candidates}
        self.assertEqual(by_text["庭审要点"].location.page_or_slide, 1)
        self.assertIn("HIDDEN_CONTENT", by_text["庭审要点"].risk_flags)
        self.assertEqual(by_text["律师备注"].kind, DocumentCandidateKind.SPEAKER_NOTE)
        self.assertIn("SPEAKER_NOTES_CONTENT", result.document_risk_flags)

    def test_rtf_plain_text_and_unicode_are_read_but_active_fields_are_blocked(self) -> None:
        safe = b"{\\rtf1\\ansi\\ansicpg1252 First paragraph\\par Unicode \\u27721?\\u35772?}"
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "note.rtf", safe, CommonDocumentFormat.RTF)
            result = read_materialized_common_document(source)
        self.assertTrue(any("First paragraph" in item.text for item in result.candidates))
        bad = b"{\\rtf1\\ansi {\\field{\\*\\fldinst INCLUDETEXT http://evil}}}"
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "bad.rtf", bad, CommonDocumentFormat.RTF)
            with self.assertRaises(CommonDocumentReadingBlocked):
                read_materialized_common_document(source)

    def test_txt_csv_html_and_eml_have_locations_and_risk_routing(self) -> None:
        samples = (
            (
                "note.txt", "第一段\n\n第二段".encode(), CommonDocumentFormat.TXT,
                lambda result: self.assertEqual(result.candidates[1].location.line_start, 3),
            ),
            (
                "rows.csv", "日期,金额\n2020-01-01,=2+2\n".encode(), CommonDocumentFormat.CSV,
                lambda result: self.assertTrue(
                    any("FORMULA_LIKE_LITERAL" in item.risk_flags for item in result.candidates)
                ),
            ),
            (
                "page.html", b"<!doctype html><html><body><p>Visible</p><a href='https://example.com'>Link</a></body></html>",
                CommonDocumentFormat.HTML,
                lambda result: self.assertIn("EXTERNAL_REFERENCE_PRESENT_NOT_FETCHED", result.document_risk_flags),
            ),
            (
                "mail.eml",
                b"From: a@example.com\r\nTo: b@example.com\r\nSubject: Case\r\nMIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=x\r\n\r\n--x\r\nContent-Type: text/plain; charset=utf-8\r\n\r\nBody\r\n--x\r\nContent-Type: application/octet-stream\r\nContent-Disposition: attachment; filename=run.exe\r\nContent-Transfer-Encoding: base64\r\n\r\nTVqQAAMAAAAEAAAA\r\n--x--\r\n",
                CommonDocumentFormat.EML,
                lambda result: self.assertIn("EXECUTABLE_ATTACHMENT_QUARANTINED", result.document_risk_flags),
            ),
        )
        for name, content, kind, assertion in samples:
            with self.subTest(kind=kind), TemporaryDirectory() as temporary:
                source = self.source(Path(temporary), name, content, kind)
                result = read_materialized_common_document(source)
                assertion(result)
                self.assertTrue(all(item.location.container_part for item in result.candidates))

    def test_attached_eml_is_not_recursively_trusted_as_inline_body(self) -> None:
        content = (
            b"From: a@example.com\r\nTo: b@example.com\r\nSubject: Outer\r\n"
            b"MIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=x\r\n\r\n"
            b"--x\r\nContent-Type: text/plain\r\n\r\nOuter body\r\n"
            b"--x\r\nContent-Type: message/rfc822\r\nContent-Disposition: attachment; filename=inner.eml\r\n\r\n"
            b"From: evil@example.com\r\nSubject: Inner\r\n\r\nINNER SECRET BODY\r\n--x--\r\n"
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "mail.eml", content, CommonDocumentFormat.EML)
            result = read_materialized_common_document(source)
        self.assertTrue(any(item.text == "inner.eml" for item in result.candidates))
        self.assertFalse(any("INNER SECRET BODY" in item.text for item in result.candidates))
        self.assertIn("ATTACHMENT_NOT_RECURSIVELY_PARSED", result.document_risk_flags)

    def test_html_script_event_handler_and_rtf_disguise_are_blocked(self) -> None:
        cases = (
            (b"<!doctype html><html><script>alert(1)</script></html>", CommonDocumentFormat.HTML),
            (b"<!doctype html><html><body onload='x()'>x</body></html>", CommonDocumentFormat.HTML),
            (b"{\\rtf1 disguised}", CommonDocumentFormat.TXT),
        )
        for index, (content, kind) in enumerate(cases):
            with self.subTest(index=index), TemporaryDirectory() as temporary:
                source = self.source(Path(temporary), "input", content, kind)
                with self.assertRaises(CommonDocumentReadingBlocked):
                    read_materialized_common_document(source)

    def test_zip_bomb_and_candidate_budget_fail_closed(self) -> None:
        bomb = self.package(
            {
                "[Content_Types].xml": "<Types/>",
                "word/document.xml": "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
                + "A" * (2 * 1024 * 1024)
                + "</w:document>",
            },
            compression=zipfile.ZIP_DEFLATED,
        )
        with TemporaryDirectory() as temporary:
            source = self.source(Path(temporary), "bomb.docx", bomb, CommonDocumentFormat.DOCX)
            with self.assertRaises(CommonDocumentReadingBlocked):
                read_materialized_common_document(
                    source,
                    budget=DocumentReadBudget(
                        max_archive_member_bytes=3 * 1024 * 1024,
                        max_archive_total_bytes=4 * 1024 * 1024,
                        max_compression_ratio=2,
                    ),
                )
        with TemporaryDirectory() as temporary:
            source = self.source(
                Path(temporary), "many.txt", b"one\n\ntwo\n\nthree", CommonDocumentFormat.TXT
            )
            with self.assertRaisesRegex(CommonDocumentReadingBlocked, "candidate budget"):
                read_materialized_common_document(
                    source, budget=DocumentReadBudget(max_candidates=2)
                )

    def test_hash_mismatch_symlink_and_repr_do_not_leak_raw_source_or_object_version(self) -> None:
        secret = "VERY-SECRET-LAWYER-TEXT"
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = self.source(root, "note.txt", secret.encode(), CommonDocumentFormat.TXT)
            self.assertNotIn(secret, repr(source))
            self.assertNotIn("version-opaque-1", repr(source))
            changed = MaterializedDocumentSource(
                source.source_object_id, "version-opaque-1", root, source.path,
                source.byte_size, "0" * 64, source.admitted_format,
            )
            with self.assertRaisesRegex(CommonDocumentReadingBlocked, "hash"):
                read_materialized_common_document(changed)
            link = root / "link.txt"
            link.symlink_to(source.path)
            linked = MaterializedDocumentSource(
                str(uuid4()), "opaque", root, link, source.byte_size,
                source.content_sha256, CommonDocumentFormat.TXT,
            )
            with self.assertRaisesRegex(CommonDocumentReadingBlocked, "symbolic"):
                read_materialized_common_document(linked)


if __name__ == "__main__":
    unittest.main()
