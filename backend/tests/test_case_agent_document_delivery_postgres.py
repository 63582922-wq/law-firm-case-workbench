from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from uuid import UUID, uuid4, uuid5
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter

from case_kernel.case_agent_document_delivery import (
    AuthoritativeDocumentSource,
    DocumentSourceKind,
    DynamicDocumentTaskBinding,
    ReviewableDocumentFormat,
    ReviewableDocumentTemplate,
    ReviewableDocumentTemplateRegistry,
    build_document_draft_request,
    build_deterministic_case_review_memo_candidate,
    build_deterministic_payment_ledger_candidate,
    canonical_document_candidate_bytes,
    first_release_reviewable_document_templates,
    parse_reviewable_document_candidate,
)
from case_kernel.case_agent_document_delivery_postgres import (
    CaseAgentDocumentPackageBlocked,
    DocumentAwareManagedArtifactAccessPort,
    PostgresReviewableDocumentPackageAccessPort,
    PostgresReviewableDocumentPackageStore,
    PrivateDocumentObjectReceipt,
    ReviewableDocumentPackageStagingRequest,
    ReviewableDocumentS3Config,
    S3ReviewableDocumentPrivateObjectStore,
    StagedReviewableDocumentPackage,
    _DOCUMENT_DELIVERY_REQUIRED_COLUMNS,
    _DOCUMENT_DELIVERY_REQUIRED_TRIGGERS,
    _INSERT_PACKAGE_SQL,
    _READ_PACKAGE_FOR_VERIFIER_SQL,
    _binding_hash,
    _artifact_ids,
    _assert_payment_ledger_matches_source_manifest,
    _assert_authorized_source_manifest_is_current,
    _authorized_source_manifest_from_sources,
    _source_set_hash_from_manifest,
    _authorized_source_refs_hash,
    _candidate_source_refs,
    _package_receipt_hash,
    _review_input_hash,
    _semantic_candidate_hash,
    _validate_staging_request,
    preflight_case_agent_document_delivery_runtime_contract,
    reviewable_document_template_hash,
)
from case_kernel.case_agent_document_adapters import ReviewableDocumentPackageStaging
from case_kernel.case_agent_verifier import (
    ManagedArtifactRead,
    first_release_review_candidate_verifiers,
)
from case_kernel.case_agent_supervisor import ArtifactReceipt
from case_kernel.approved_draft_worker import DraftArtifact
from case_kernel.office_pdf_conversion_worker import ConvertedOfficePdf
from case_kernel.reviewable_draft_worker import ReviewableOfficeDraft
from case_kernel.case_work_plan import (
    CaseWorkPlanItem,
    DeliveryTarget,
    ReviewGate,
    WorkPlanItemKind,
    WorkPlanReadiness,
)
from case_kernel.models import Actor, Role
from case_kernel.web_object_store import (
    S3CompatiblePrivateObjectStore,
    S3PrivateObjectStoreConfig,
)
from backend.tests.test_case_agent_document_delivery import (
    lawyer_decision_package_payload,
)


def digest(value: bytes | str) -> str:
    raw = value if isinstance(value, bytes) else value.encode()
    return sha256(raw).hexdigest()


def canonical(value: object) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode()


def docx_bytes() -> bytes:
    value = Document()
    value.add_heading("案件审阅意见候选", 0)
    value.add_paragraph("仅供律师复核。")
    output = BytesIO()
    value.save(output)
    return output.getvalue()


def pdf_bytes() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(output)
    return output.getvalue()


class _Cursor:
    def __init__(self, row=None):
        self.row = row

    def fetchone(self):
        return self.row

    def fetchall(self):
        return self.row or []


class _Connection:
    def __init__(self, rows):
        self.rows = list(rows)
        self.sql = []

    def execute(self, sql, params=()):
        self.sql.append((" ".join(sql.split()), params))
        if "SELECT package.*" in sql:
            return _Cursor(self.rows.pop(0) if self.rows else None)
        return _Cursor()


class _Context:
    def __init__(self, connection): self.connection = connection
    def __enter__(self): return self.connection
    def __exit__(self, *_): return False


class _Objects:
    def __init__(self):
        self.values = {}

    def put_reviewable_document_object(
        self, content, *, firm_id, matter_id, package_id, object_role,
        content_sha256, media_type,
    ):
        key = (
            f"case-agent-document-packages/v1/{firm_id}/{matter_id}/"
            f"{package_id}/{object_role}/{content_sha256}.lca"
        )
        receipt = PrivateDocumentObjectReceipt(
            key, content_sha256, len(content), media_type, "v1"
        )
        self.values[key] = content
        return receipt

    def read_reviewable_document_object(self, receipt):
        return self.values[receipt.object_key]


class _PreflightConnection:
    def __init__(
        self,
        *,
        omit_table=False,
        omit_column=None,
        omit_trigger=None,
        force_rls=True,
    ):
        self.omit_table = omit_table
        self.omit_column = omit_column
        self.omit_trigger = omit_trigger
        self.force_rls = force_rls
        self.executions = []

    def execute(self, sql, params=()):
        normalized = " ".join(sql.split())
        self.executions.append((normalized, params))
        if "information_schema.columns" in normalized:
            rows = [] if self.omit_table else [
                {
                    "table_name": "case_agent_reviewable_document_packages",
                    "column_name": column,
                }
                for column in sorted(_DOCUMENT_DELIVERY_REQUIRED_COLUMNS)
                if column != self.omit_column
            ]
            return _Cursor(rows)
        if "information_schema.triggers" in normalized:
            return _Cursor(
                [
                    {"trigger_name": trigger}
                    for trigger in sorted(_DOCUMENT_DELIVERY_REQUIRED_TRIGGERS)
                    if trigger != self.omit_trigger
                ]
            )
        if "pg_catalog.pg_class" in normalized:
            if self.omit_table:
                return _Cursor(None)
            return _Cursor(
                {
                    "relrowsecurity": self.force_rls,
                    "relforcerowsecurity": self.force_rls,
                }
            )
        raise AssertionError(f"unexpected preflight SQL: {normalized}")


class _S3:
    def __init__(self):
        self.values = {}

    def put_object(self, **kwargs):
        item = dict(kwargs)
        body = item.pop("Body")
        content = body.read() if callable(getattr(body, "read", None)) else bytes(body)
        self.values[(kwargs["Bucket"], kwargs["Key"])] = {
            **item,
            "content": content,
        }
        return {"VersionId": "v1"}

    def head_object(self, **kwargs):
        return self.values[(kwargs["Bucket"], kwargs["Key"])]

    def get_object(self, **kwargs):
        value = self.values[(kwargs["Bucket"], kwargs["Key"])]
        return {"Body": BytesIO(value["content"])}

    def delete_object(self, **kwargs):
        self.values.pop((kwargs["Bucket"], kwargs["Key"]), None)
        return {}


class CaseAgentDocumentDeliveryPostgresTests(unittest.TestCase):
    def setUp(self):
        self.firm = str(uuid4())
        self.matter = str(uuid4())
        self.run = str(uuid4())
        self.graph = str(uuid4())
        self.task = str(uuid4())
        self.attempt = str(uuid4())
        self.plan = str(uuid4())
        self.item = str(uuid4())
        self.profile = str(uuid4())
        self.fact = str(uuid4())
        self.decision_package = str(uuid4())
        self.task_hash = digest("task")
        self.snapshot_hash = digest("snapshot")
        self.plan_hash = digest("plan")
        self.profile_hash = digest("profile")
        self.verification_hash = digest("lawyer-verification")
        item = CaseWorkPlanItem(
            item_id=self.item,
            sequence=1,
            kind=WorkPlanItemKind.DOCUMENT_CANDIDATE,
            readiness=WorkPlanReadiness.ACTIONABLE,
            title="形成案件审阅候选",
            purpose="将已确认的案件范围整理为可复核的内部文书。",
            rationale="已具备当前态势和工作计划。",
            prerequisites=(),
            risk_if_omitted="无法集中核对已确认内容。",
            confidence=0.9,
            review_gate=ReviewGate.LEAD_LAWYER_CONFIRMATION,
            delivery_target=DeliveryTarget.INTERNAL_WORK_PRODUCT,
            deliverable_kind="CASE_REVIEW_MEMO",
            required_for_delivery=False,
            is_primary_document=False,
            trigger_refs=(),
            source_refs=(),
        )
        self.template = first_release_reviewable_document_templates().get(
            "CASE_REVIEW_MEMO"
        )
        decision_package_payload = lawyer_decision_package_payload()
        self.sources = (
            AuthoritativeDocumentSource(
                input_ref=f"posture-profile:{self.profile}",
                source_kind=DocumentSourceKind.POSTURE_PROFILE,
                source_version="v1",
                source_hash=self.profile_hash,
                label="当前代理态势",
                text=json.dumps(
                    {
                        "represented_party": "测试公司",
                        "represented_position": "PLAINTIFF",
                        "procedure_stage": "PRE_ACTION",
                        "case_type_code": "SALE_CONTRACT_DISPUTE",
                        "authority_scope_code": "LITIGATION_FULL",
                        "engagement_state": "ACTIVE",
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ),
            ),
            AuthoritativeDocumentSource(
                input_ref=f"work-plan-item:{self.item}",
                source_kind=DocumentSourceKind.WORK_PLAN_ITEM,
                source_version="v1",
                source_hash=self.plan_hash,
                label="当前办案计划",
                text=json.dumps(
                    {
                        "title": item.title,
                        "purpose": item.purpose,
                        "rationale": item.rationale,
                        "risk_if_omitted": item.risk_if_omitted,
                        "delivery_target": item.delivery_target.value,
                        "deliverable_kind": item.deliverable_kind,
                        "required_for_delivery": item.required_for_delivery,
                        "is_primary_document": item.is_primary_document,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ),
            ),
            AuthoritativeDocumentSource(
                input_ref=f"fact:{self.fact}",
                source_kind=DocumentSourceKind.CONFIRMED_FACT,
                source_version="v1",
                source_hash=digest("fact-decision"),
                label="已确认事实",
                text="借款已经实际交付。",
            ),
            AuthoritativeDocumentSource(
                input_ref=f"lawyer-decision-package:{self.decision_package}",
                source_kind=(
                    DocumentSourceKind.VERIFIED_LAWYER_DECISION_PACKAGE
                ),
                source_version=f"verified-{self.verification_hash}",
                source_hash=digest(decision_package_payload),
                label="独立校验通过的律师决策包",
                text=decision_package_payload.decode("utf-8"),
            ),
        )
        self.binding = DynamicDocumentTaskBinding(
            firm_id=self.firm,
            matter_id=self.matter,
            run_id=self.run,
            graph_id=self.graph,
            task_id=self.task,
            task_input_hash=self.task_hash,
            case_snapshot_hash=self.snapshot_hash,
            work_plan_id=self.plan,
            work_plan_hash=self.plan_hash,
            work_plan_status="ACTIVE",
            work_plan_item=item,
            posture_profile_id=self.profile,
            posture_profile_hash=self.profile_hash,
            template=self.template,
            sources=self.sources,
        )
        self.candidate = build_deterministic_case_review_memo_candidate(
            self.binding
        )
        self.candidate_bytes = canonical_document_candidate_bytes(self.candidate)
        self.editable = docx_bytes()
        self.pdf = pdf_bytes()
        preliminary = ReviewableDocumentPackageStagingRequest(
            idempotency_key=digest("idempotency"),
            run_id=self.run,
            graph_id=self.graph,
            task_id=self.task,
            attempt_id=self.attempt,
            task_input_hash=self.task_hash,
            case_snapshot_hash=self.snapshot_hash,
            binding_hash=self.binding.binding_hash,
            source_set_hash=self.binding.source_set_hash,
            authorized_source_refs=tuple(
                sorted(item.input_ref for item in self.sources)
            ),
            authorized_source_manifest=_authorized_source_manifest_from_sources(
                self.sources
            ),
            candidate_hash=self.candidate.candidate_hash,
            work_plan_id=self.plan,
            work_plan_hash=self.plan_hash,
            work_plan_item_id=self.item,
            posture_profile_id=self.profile,
            posture_profile_hash=self.profile_hash,
            template_id=self.template.template_id,
            template_version=self.template.template_version,
            template_hash=reviewable_document_template_hash(self.template),
            deliverable_kind=self.template.deliverable_kind,
            output_format=ReviewableDocumentFormat.DOCX,
            candidate_bytes=self.candidate_bytes,
            candidate_content_sha256=digest(self.candidate_bytes),
            editable_bytes=self.editable,
            editable_sha256=digest(self.editable),
            editable_media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            review_pdf_bytes=self.pdf,
            review_pdf_sha256=digest(self.pdf),
            review_pdf_page_count=1,
            render_verification_hash=digest("render"),
            review_input_hash=digest("placeholder"),
        )
        self.request = replace(
            preliminary, review_input_hash=_review_input_hash(preliminary)
        )

    def test_content_package_contract_is_bound_but_cannot_write_before_registration_exists(self):
        from unittest.mock import Mock
        import case_kernel.case_agent_document_delivery_postgres as delivery

        request = replace(self.request, generation_mode="LAWYER_CONTENT_REVISION", revision_number=2,
            content_generation_claim_version=2,
            root_package_id=str(uuid4()), supersedes_package_id=str(uuid4()),
            revision_request_id=str(uuid4()), requested_by=str(uuid4()))
        _validate_staging_request(request, first_release_reviewable_document_templates())
        for field, value in (("revision_number", True), ("revision_number", 2.0), ("revision_number", 1),
                ("revision_request_id", None), ("requested_by", None), ("root_package_id", None),
                ("supersedes_package_id", None), ("output_format", ReviewableDocumentFormat.XLSX),
                ("content_generation_claim_version", None), ("content_generation_claim_version", True),
                ("content_generation_claim_version", 0), ("content_generation_claim_version", 4)):
            with self.subTest(field=field, value=value), self.assertRaises(CaseAgentDocumentPackageBlocked):
                _validate_staging_request(replace(request, **{field: value}), first_release_reviewable_document_templates())
        worker = Actor(str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER}))
        objects = _Objects()
        store = PostgresReviewableDocumentPackageStore(dsn="test-only", worker_actor=worker, object_store=objects)
        with patch.object(delivery, "_transaction", Mock(side_effect=AssertionError("database must not be opened"))) as transaction:
            with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "registration is not enabled"):
                store.stage_package(request)
            transaction.assert_not_called()
        self.assertEqual(objects.values, {})

    def test_package_receipt_distinguishes_content_mode_and_each_revision_coordinate(self):
        import case_kernel.case_agent_document_delivery_postgres as delivery

        package_id = str(uuid4())
        args = dict(package_id=package_id, artifact_ids=_artifact_ids(self.task, package_id),
            template_hash=self.request.template_hash,
            object_receipts={role: PrivateDocumentObjectReceipt(f"private/{role}", digest(role), 12, media, "v1")
                for role, media in (("candidate", "application/json"), ("editable", self.request.editable_media_type),
                                    ("pdf-preview", "application/pdf"))})
        revision = replace(self.request, generation_mode="DETERMINISTIC_TEMPLATE_REVISION", revision_number=2,
            root_package_id=str(uuid4()), supersedes_package_id=str(uuid4()),
            revision_request_id=str(uuid4()), requested_by=str(uuid4()))
        content = replace(revision, generation_mode="LAWYER_CONTENT_REVISION", content_generation_claim_version=2)
        with self.assertRaises(CaseAgentDocumentPackageBlocked):
            _validate_staging_request(replace(revision, content_generation_claim_version=2),
                first_release_reviewable_document_templates())
        initial_hash = _package_receipt_hash(request=self.request, **args)
        template_hash = _package_receipt_hash(request=revision, **args)
        content_hash = _package_receipt_hash(request=content, **args)
        self.assertEqual(len({initial_hash, template_hash, content_hash}), 3)
        for field, value in (("revision_number", 3), ("root_package_id", str(uuid4())),
                ("supersedes_package_id", str(uuid4())), ("revision_request_id", str(uuid4())),
                ("requested_by", str(uuid4())), ("content_generation_claim_version", 3)):
            with self.subTest(field=field):
                self.assertNotEqual(content_hash, _package_receipt_hash(request=replace(content, **{field: value}), **args))
        with patch.object(delivery, "_canonical_hash", side_effect=lambda payload: payload):
            initial = _package_receipt_hash(request=self.request, **args)
            template = _package_receipt_hash(request=revision, **args)
            body = _package_receipt_hash(request=content, **args)
        self.assertEqual(initial["schema_version"], "case-agent-reviewable-document-package-receipt-v1")
        self.assertNotIn("generation_mode", initial)
        self.assertEqual(template["schema_version"], "case-agent-reviewable-document-package-receipt-v2")
        self.assertEqual(body, {**template, "generation_mode": "LAWYER_CONTENT_REVISION", "content_generation_claim_version": 2})
        self.assertEqual(body["review_status"], "NEEDS_LAWYER_REVIEW")
        self.assertFalse(body["formal_document"])
        self.assertFalse(body["court_submitted"])
        with self.assertRaises(CaseAgentDocumentPackageBlocked):
            _package_receipt_hash(request=replace(content, generation_mode="UNKNOWN"), **args)

    def test_core_and_postgres_candidate_hashes_are_identical(self):
        value = json.loads(self.candidate_bytes)
        self.assertEqual(_semantic_candidate_hash(value), self.candidate.candidate_hash)
        normalized = _validate_staging_request(
            self.request, first_release_reviewable_document_templates()
        )
        self.assertEqual(normalized["candidate"], value)
        self.assertEqual(
            _candidate_source_refs(value),
            frozenset(item.input_ref for item in self.sources),
        )

    def test_package_verifier_matches_each_payment_row_to_source_text_digest(self):
        from backend.tests.test_case_agent_document_delivery import (
            CaseAgentDocumentDeliveryTests,
        )

        binding = CaseAgentDocumentDeliveryTests().binding("PAYMENT_LEDGER")
        candidate = build_deterministic_payment_ledger_candidate(binding)
        value = json.loads(canonical_document_candidate_bytes(candidate))
        manifest = _authorized_source_manifest_from_sources(binding.sources)
        _assert_payment_ledger_matches_source_manifest(
            candidate=value,
            deliverable_kind="PAYMENT_LEDGER",
            output_format=ReviewableDocumentFormat.XLSX,
            manifest=manifest,
        )
        value["rows"][0]["cells"]["amount"] = "999.000000"
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "values differ"
        ):
            _assert_payment_ledger_matches_source_manifest(
                candidate=value,
                deliverable_kind="PAYMENT_LEDGER",
                output_format=ReviewableDocumentFormat.XLSX,
                manifest=manifest,
            )

    def test_server_template_hash_detects_same_version_instruction_drift(self):
        self.assertEqual(
            reviewable_document_template_hash(self.template),
            self.template.template_hash,
        )
        changed = replace(
            self.template,
            drafting_instructions=self.template.drafting_instructions + ("不同指令",),
        )
        self.assertNotEqual(
            reviewable_document_template_hash(self.template),
            reviewable_document_template_hash(changed),
        )
        changed_registry = ReviewableDocumentTemplateRegistry((changed,))
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "server template"
        ):
            _validate_staging_request(self.request, changed_registry)

    def test_explicit_template_hash_must_match_the_server_registry(self):
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "server template"
        ):
            _validate_staging_request(
                replace(self.request, template_hash=digest("wrong template")),
                first_release_reviewable_document_templates(),
            )

    def test_binding_hash_requires_real_tenant_coordinates(self):
        self.assertEqual(
            _binding_hash(self.request, firm_id=self.firm, matter_id=self.matter),
            self.binding.binding_hash,
        )

    def test_artifact_ids_match_the_worker_adapter_deterministic_contract(self):
        package_id = str(uuid5(UUID(self.task), self.request.idempotency_key))
        artifact_ids = _artifact_ids(self.task, package_id)
        for role, kind in (
            ("candidate", "REVIEWABLE_DOCUMENT_CANDIDATE_JSON"),
            ("editable", "REVIEWABLE_DOCUMENT_EDITABLE"),
            ("pdf-preview", "REVIEWABLE_DOCUMENT_PDF_PREVIEW"),
        ):
            self.assertEqual(
                artifact_ids[role],
                str(uuid5(UUID(self.task), f"{package_id}:{kind}")),
            )

    def test_authorized_source_refs_are_separate_and_canonical(self):
        refs = self.request.authorized_source_refs
        self.assertEqual(refs, tuple(sorted(item.input_ref for item in self.sources)))
        self.assertEqual(len(_authorized_source_refs_hash(refs)), 64)
        self.assertEqual(
            _source_set_hash_from_manifest(self.request.authorized_source_manifest),
            self.request.source_set_hash,
        )
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "sorted"):
            _validate_staging_request(
                replace(self.request, authorized_source_refs=tuple(reversed(refs))),
                first_release_reviewable_document_templates(),
            )
        self.assertNotEqual(
            _binding_hash(
                self.request, firm_id=str(uuid4()), matter_id=self.matter
            ),
            self.binding.binding_hash,
        )

    def test_source_manifest_rejects_ref_or_text_digest_drift(self):
        manifest = self.request.authorized_source_manifest
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "manifest"):
            _validate_staging_request(
                replace(
                    self.request,
                    authorized_source_manifest=(
                        replace(manifest[0], text_sha256=digest("different text")),
                        *manifest[1:],
                    ),
                ),
                first_release_reviewable_document_templates(),
            )
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "refs differ"):
            _validate_staging_request(
                replace(
                    self.request,
                    authorized_source_manifest=(manifest[0],),
                ),
                first_release_reviewable_document_templates(),
            )

    def test_source_manifest_is_rechecked_against_current_plan_and_posture(self):
        current_rows = [
            {
                "source_type": "CASE_FACT",
                "source_id": self.fact,
                "source_version": "v1",
                "source_hash": digest("fact-decision"),
                "reference_use": "FACT",
            }
        ]

        class CurrentSources:
            def execute(_self, sql, _params=()):
                normalized = " ".join(sql.split())
                if "FROM case_work_plan_item_references" in normalized:
                    return _Cursor(current_rows)
                if "FROM case_agent_artifacts artifact" in normalized:
                    return _Cursor(
                        [
                            {
                                "content_hash": digest(
                                    lawyer_decision_package_payload()
                                ),
                                "verification_hash": self.verification_hash,
                            }
                        ]
                    )
                raise AssertionError(
                    f"unexpected current-source SQL: {normalized}"
                )

        _assert_authorized_source_manifest_is_current(
            CurrentSources(),
            firm_id=self.firm,
            matter_id=self.matter,
            work_plan_id=self.plan,
            work_plan_item_id=self.item,
            work_plan_version=1,
            work_plan_hash=self.plan_hash,
            posture_profile_id=self.profile,
            posture_profile_version=1,
            posture_profile_hash=self.profile_hash,
            manifest=self.request.authorized_source_manifest,
        )
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "current plan"):
            _assert_authorized_source_manifest_is_current(
                CurrentSources(),
                firm_id=self.firm,
                matter_id=self.matter,
                work_plan_id=self.plan,
                work_plan_item_id=self.item,
                work_plan_version=2,
                work_plan_hash=self.plan_hash,
                posture_profile_id=self.profile,
                posture_profile_version=1,
                posture_profile_hash=self.profile_hash,
                manifest=self.request.authorized_source_manifest,
            )

        class StagingSources:
            def execute(_self, sql, _params=()):
                normalized = " ".join(sql.split())
                if "FROM case_work_plan_item_references" in normalized:
                    return _Cursor(current_rows)
                raise AssertionError(
                    "execution staging must not read verifier-owned history"
                )

        _assert_authorized_source_manifest_is_current(
            StagingSources(),
            firm_id=self.firm,
            matter_id=self.matter,
            work_plan_id=self.plan,
            work_plan_item_id=self.item,
            work_plan_version=1,
            work_plan_hash=self.plan_hash,
            posture_profile_id=self.profile,
            posture_profile_version=1,
            posture_profile_hash=self.profile_hash,
            manifest=self.request.authorized_source_manifest,
            independent_artifact_recheck=False,
        )

    def test_promoted_fact_and_transaction_bindings_are_unwrapped_to_current_sources(self):
        transaction_id = str(uuid4())
        transaction_hash = digest("transaction-confirmation")
        transaction_source = AuthoritativeDocumentSource(
            input_ref=f"transaction:{transaction_id}",
            source_kind=DocumentSourceKind.CONFIRMED_TRANSACTION,
            source_version="v1",
            source_hash=transaction_hash,
            label="已确认付款",
            text=json.dumps(
                {
                    "local_date": "2025-04-10",
                    "amount": "30000.000000",
                    "currency": "CNY",
                    "transaction_reference": "TEST-BANK-003000",
                },
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
        )
        cases = (
            (
                "fact",
                "FACT",
                self.request.authorized_source_manifest,
                {
                    "object_type": "CASE_FACT",
                    "object_id": self.fact,
                    "object_version": "v1",
                    "source_status": "CONFIRMED",
                    "reference_use": "FACT",
                    "snapshot_matter_version": 1,
                    "fact_id": self.fact,
                    "fact_status": "CONFIRMED",
                    "decision_hash": digest("fact-decision"),
                    "transaction_id": None,
                    "transaction_status": None,
                    "confirmation_hash": None,
                },
            ),
            (
                "transaction",
                "TRANSACTION",
                _authorized_source_manifest_from_sources(
                    (self.sources[0], self.sources[1], transaction_source)
                ),
                {
                    "object_type": "CASE_TRANSACTION",
                    "object_id": transaction_id,
                    "object_version": "v1",
                    "source_status": "CONFIRMED",
                    "reference_use": "TRANSACTION",
                    "snapshot_matter_version": 1,
                    "fact_id": None,
                    "fact_status": None,
                    "decision_hash": None,
                    "transaction_id": transaction_id,
                    "transaction_status": "CONFIRMED",
                    "confirmation_hash": transaction_hash,
                },
            ),
        )
        for label, reference_use, manifest, binding_row in cases:
            binding_id = str(uuid4())
            reference_rows = [
                {
                    "source_type": "AGENT_TASK_INPUT",
                    "source_id": binding_id,
                    "source_version": "v1",
                    "source_hash": digest(f"{label}-binding"),
                    "reference_use": reference_use,
                }
            ]

            class CurrentPromotedSource:
                def execute(_self, sql, _params=()):
                    normalized = " ".join(sql.split())
                    if "FROM case_work_plan_item_references" in normalized:
                        return _Cursor(reference_rows)
                    if "FROM case_agent_work_plan_input_bindings" in normalized:
                        return _Cursor(binding_row)
                    if "FROM case_agent_artifacts artifact" in normalized:
                        return _Cursor(
                            [
                                {
                                    "content_hash": digest(
                                        lawyer_decision_package_payload()
                                    ),
                                    "verification_hash": self.verification_hash,
                                }
                            ]
                        )
                    raise AssertionError(f"unexpected current-source SQL: {normalized}")

            with self.subTest(source=label):
                _assert_authorized_source_manifest_is_current(
                    CurrentPromotedSource(),
                    firm_id=self.firm,
                    matter_id=self.matter,
                    work_plan_id=self.plan,
                    work_plan_item_id=self.item,
                    work_plan_version=1,
                    work_plan_hash=self.plan_hash,
                    posture_profile_id=self.profile,
                    posture_profile_version=1,
                    posture_profile_hash=self.profile_hash,
                    manifest=manifest,
                )

    def test_promoted_source_rejects_unsupported_or_changed_ledger_identity(self):
        binding_id = str(uuid4())
        reference_rows = [
            {
                "source_type": "AGENT_TASK_INPUT",
                "source_id": binding_id,
                "source_version": "v1",
                "source_hash": digest("fact-binding"),
                "reference_use": "FACT",
            }
        ]
        current_fact = {
            "object_type": "CASE_FACT",
            "object_id": self.fact,
            "object_version": "v1",
            "source_status": "CONFIRMED",
            "reference_use": "FACT",
            "snapshot_matter_version": 1,
            "fact_id": self.fact,
            "fact_status": "CONFIRMED",
            "decision_hash": digest("changed-decision"),
            "transaction_id": None,
            "transaction_status": None,
            "confirmation_hash": None,
        }

        class CurrentPromotedSource:
            binding_row = current_fact

            def execute(_self, sql, _params=()):
                normalized = " ".join(sql.split())
                if "FROM case_work_plan_item_references" in normalized:
                    return _Cursor(reference_rows)
                if "FROM case_agent_work_plan_input_bindings" in normalized:
                    return _Cursor(_self.binding_row)
                raise AssertionError(f"unexpected current-source SQL: {normalized}")

        current = CurrentPromotedSource()
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "current plan"):
            _assert_authorized_source_manifest_is_current(
                current,
                firm_id=self.firm,
                matter_id=self.matter,
                work_plan_id=self.plan,
                work_plan_item_id=self.item,
                work_plan_version=1,
                work_plan_hash=self.plan_hash,
                posture_profile_id=self.profile,
                posture_profile_version=1,
                posture_profile_hash=self.profile_hash,
                manifest=self.request.authorized_source_manifest,
            )
        current.binding_row = {
            **current_fact,
            "object_type": "EVIDENCE_PAGE",
            "reference_use": "EVIDENCE",
        }
        reference_rows[0] = {**reference_rows[0], "reference_use": "EVIDENCE"}
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "not approved for disclosure"
        ):
            _assert_authorized_source_manifest_is_current(
                current,
                firm_id=self.firm,
                matter_id=self.matter,
                work_plan_id=self.plan,
                work_plan_item_id=self.item,
                work_plan_version=1,
                work_plan_hash=self.plan_hash,
                posture_profile_id=self.profile,
                posture_profile_version=1,
                posture_profile_hash=self.profile_hash,
                manifest=self.request.authorized_source_manifest,
            )

    def test_private_object_locator_never_appears_in_receipt_repr(self):
        receipt = PrivateDocumentObjectReceipt(
            object_key="case-agent-document-packages/v1/private",
            content_sha256=digest("x"),
            byte_size=1,
            media_type="application/json",
        )
        self.assertNotIn("case-agent-document-packages", repr(receipt))

    def test_s3_bridge_uses_exact_private_0039_key_and_reads_back(self):
        client = _S3()
        store = S3ReviewableDocumentPrivateObjectStore(
            config=ReviewableDocumentS3Config(bucket="lawcase-private"),
            client=client,
        )
        package_id = str(uuid4())
        receipt = store.put_reviewable_document_object(
            self.candidate_bytes,
            firm_id=self.firm,
            matter_id=self.matter,
            package_id=package_id,
            object_role="candidate",
            content_sha256=digest(self.candidate_bytes),
            media_type="application/json",
        )
        self.assertTrue(
            receipt.object_key.startswith(
                f"case-agent-document-packages/v1/{self.firm}/{self.matter}/"
                f"{package_id}/candidate/"
            )
        )
        self.assertEqual(
            store.read_reviewable_document_object(receipt), self.candidate_bytes
        )

    def test_existing_web_private_store_implements_the_document_object_bridge(self):
        client = _S3()
        store = S3CompatiblePrivateObjectStore(
            S3PrivateObjectStoreConfig(
                endpoint_url="http://object-storage:9000",
                region_name="us-east-1",
                bucket="lawcase-private",
                access_key_id="lawcase-app-user",
                secret_access_key="x" * 32,
                allow_insecure_internal_endpoint=True,
            ),
            client=client,
        )
        package_id = str(uuid4())
        receipt = store.put_reviewable_document_object(
            self.candidate_bytes,
            firm_id=self.firm,
            matter_id=self.matter,
            package_id=package_id,
            object_role="candidate",
            content_sha256=digest(self.candidate_bytes),
            media_type="application/json",
        )
        self.assertNotIn(receipt.object_key, repr(receipt))
        self.assertEqual(
            store.read_reviewable_document_object(receipt), self.candidate_bytes
        )

    def test_worker_staging_dto_bridge_expands_the_authorized_manifest(self):
        actor = Actor(str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER}))
        store = PostgresReviewableDocumentPackageStore(
            dsn="postgresql://server-owned",
            worker_actor=actor,
            object_store=_Objects(),
        )
        editable = DraftArtifact(
            media_type=self.request.editable_media_type,
            content=self.editable,
            content_sha256=self.request.editable_sha256,
        )
        review_pdf = ConvertedOfficePdf(
            source_sha256=self.request.editable_sha256,
            detected_kind="WORD_DOCUMENT",
            converter_id="test-isolated-converter",
            converter_version="1.0.0",
            transform_hash=digest("transform"),
            pdf_sha256=self.request.review_pdf_sha256,
            pdf_bytes=len(self.pdf),
            page_count=1,
            render_verification_hash=self.request.render_verification_hash,
            pdf_content=self.pdf,
        )
        generated = ReviewableOfficeDraft(
            editable_artifact=editable,
            review_pdf=review_pdf,
            approval_hash=self.candidate.candidate_hash,
            review_input_hash=self.request.review_input_hash,
        )
        adapter_request = ReviewableDocumentPackageStaging(
            run_id=self.run,
            task_id=self.task,
            attempt_id=self.attempt,
            task_input_hash=self.task_hash,
            binding=self.binding,
            candidate=self.candidate,
            candidate_content=self.candidate_bytes,
            generated=generated,
        )
        package_id = str(uuid5(UUID(self.task), self.candidate.candidate_hash))
        artifact_ids = _artifact_ids(self.task, package_id)
        staged = StagedReviewableDocumentPackage(
            package_id=package_id,
            candidate_artifact=ArtifactReceipt(
                artifact_ids["candidate"],
                "REVIEWABLE_DOCUMENT_CANDIDATE_JSON",
                self.request.candidate_content_sha256,
                len(self.candidate_bytes),
                self.task_hash,
                True,
            ),
            editable_artifact=ArtifactReceipt(
                artifact_ids["editable"],
                "REVIEWABLE_DOCUMENT_EDITABLE",
                self.request.editable_sha256,
                len(self.editable),
                self.task_hash,
                True,
            ),
            review_pdf_artifact=ArtifactReceipt(
                artifact_ids["pdf-preview"],
                "REVIEWABLE_DOCUMENT_PDF_PREVIEW",
                self.request.review_pdf_sha256,
                len(self.pdf),
                self.task_hash,
                True,
            ),
            receipt_hash=digest("package receipt"),
        )
        with patch.object(store, "stage_package", return_value=staged) as stage:
            result = store.stage_document_package(adapter_request)
        normalized = stage.call_args.args[0]
        self.assertEqual(normalized.authorized_source_refs, self.request.authorized_source_refs)
        self.assertEqual(
            normalized.authorized_source_manifest,
            self.request.authorized_source_manifest,
        )
        self.assertEqual(normalized.source_set_hash, self.binding.source_set_hash)
        self.assertEqual(result.artifact_receipts, staged.artifacts)

    def test_request_rejects_formula_macro_like_or_unbound_content(self):
        tampered = replace(
            self.request,
            candidate_bytes=self.candidate_bytes + b" ",
            candidate_content_sha256=digest(self.candidate_bytes + b" "),
        )
        with self.assertRaisesRegex(CaseAgentDocumentPackageBlocked, "canonical"):
            _validate_staging_request(
                tampered, first_release_reviewable_document_templates()
            )

    def test_verifier_query_reads_any_artifact_as_one_current_package(self):
        self.assertIn("SELECT package.*, task.input_refs", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("package.candidate_artifact_id = %s", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("package.editable_artifact_id = %s", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("package.review_pdf_artifact_id = %s", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("run.current_graph_id = package.graph_id", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("task.input_refs = jsonb_build_array", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("plan_head.current_plan_id = plan.plan_id", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertIn("profile_head.current_profile_id = profile.profile_id", _READ_PACKAGE_FOR_VERIFIER_SQL)
        self.assertEqual(_READ_PACKAGE_FOR_VERIFIER_SQL.count("%s"), 11)

    def test_insert_contract_has_one_parameter_per_placeholder(self):
        from case_kernel.case_agent_document_delivery_postgres import _insert_parameters

        package_id = str(uuid5(UUID(self.task), self.request.idempotency_key))
        object_receipts = {
            "candidate": PrivateDocumentObjectReceipt(
                f"case-agent-document-packages/v1/{self.firm}/{self.matter}/"
                f"{package_id}/candidate/{self.request.candidate_content_sha256}.lca",
                self.request.candidate_content_sha256,
                len(self.request.candidate_bytes),
                "application/json",
            ),
            "editable": PrivateDocumentObjectReceipt(
                f"case-agent-document-packages/v1/{self.firm}/{self.matter}/"
                f"{package_id}/editable/{self.request.editable_sha256}.lca",
                self.request.editable_sha256,
                len(self.request.editable_bytes),
                self.request.editable_media_type,
            ),
            "pdf-preview": PrivateDocumentObjectReceipt(
                f"case-agent-document-packages/v1/{self.firm}/{self.matter}/"
                f"{package_id}/pdf-preview/{self.request.review_pdf_sha256}.lca",
                self.request.review_pdf_sha256,
                len(self.request.review_pdf_bytes),
                "application/pdf",
            ),
        }
        parameters = _insert_parameters(
            request=self.request,
            firm_id=self.firm,
            matter_id=self.matter,
            staged_by=str(uuid4()),
            package_id=package_id,
            artifact_ids=_artifact_ids(self.task, package_id),
            template_hash=self.request.template_hash,
            object_receipts=object_receipts,
            package_receipt_hash=digest("package receipt"),
        )
        self.assertEqual(_INSERT_PACKAGE_SQL.count("%s"), len(parameters))
        from case_kernel.case_agent_document_delivery_postgres import _insert_package_sql
        self.assertEqual(_insert_package_sql("INITIAL_AGENT_TASK"), _INSERT_PACKAGE_SQL)
        self.assertEqual(_insert_package_sql("DETERMINISTIC_TEMPLATE_REVISION"), _INSERT_PACKAGE_SQL)
        content = replace(self.request, generation_mode="LAWYER_CONTENT_REVISION", revision_number=2,
            root_package_id=str(uuid4()), supersedes_package_id=str(uuid4()), revision_request_id=str(uuid4()),
            requested_by=str(uuid4()), content_generation_claim_version=2)
        content_parameters = _insert_parameters(request=content, firm_id=self.firm, matter_id=self.matter,
            staged_by=str(uuid4()), package_id=package_id, artifact_ids=_artifact_ids(self.task, package_id),
            template_hash=content.template_hash, object_receipts=object_receipts, package_receipt_hash=digest("package receipt"))
        content_sql = _insert_package_sql(content.generation_mode)
        self.assertEqual(content_sql.count("%s"), len(content_parameters))
        self.assertEqual(len(content_parameters), len(parameters) + 1)
        self.assertEqual(content_parameters[-1], 2)
        self.assertEqual(content_sql.count("content_generation_claim_version"), 1)
        self.assertNotIn("content_generation_claim_version", _INSERT_PACKAGE_SQL)
        with self.assertRaises(CaseAgentDocumentPackageBlocked):
            _insert_package_sql("UNKNOWN")

    def test_content_staging_binding_uses_render_job_and_exact_authorization(self):
        from unittest.mock import MagicMock
        import case_kernel.case_agent_document_delivery_postgres as delivery

        worker = Actor(str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER}))
        request = replace(self.request, generation_mode="LAWYER_CONTENT_REVISION", revision_number=2,
            root_package_id=str(uuid4()), supersedes_package_id=str(uuid4()), revision_request_id=str(uuid4()),
            requested_by=str(uuid4()), content_generation_claim_version=2)
        row = dict(matter_id=self.matter, graph_hash=digest("graph"), current_graph_hash=digest("graph"),
            snapshot_hash=request.case_snapshot_hash, graph_snapshot_hash=request.case_snapshot_hash,
            input_hash=request.task_input_hash, current_graph_id=request.graph_id, run_status="READY_FOR_REVIEW",
            is_stale=False, is_cancelled=False, matter_version=1, snapshot_matter_version=1,
            writes_managed_derivatives=True, approval_gate="LAWYER_REVIEW", head_status="SUCCEEDED", is_current=True,
            attempt_status="SUCCEEDED", worker_status="ACTIVE", worker_only=True, active_role_count=1,
            inbox_state="RENDERING", claimed_by=worker.actor_id, lease_expires_at="database-checked-live-lease",
            predecessor_package_id=request.supersedes_package_id, requested_root_package_id=request.root_package_id,
            predecessor_root_package_id=request.root_package_id, predecessor_revision_number=1, expected_revision_number=1,
            source_package_receipt_hash=digest("predecessor"), package_receipt_hash=digest("predecessor"),
            revision_requested_by=request.requested_by, target_template_id=request.template_id,
            target_template_version=request.template_version, target_template_hash=request.template_hash,
            plan_status="ACTIVE", current_plan_id=request.work_plan_id, plan_hash=request.work_plan_hash,
            profile_id=request.posture_profile_id, profile_hash=request.posture_profile_hash,
            actual_profile_hash=request.posture_profile_hash, current_profile_id=request.posture_profile_id,
            activated_matter_version=1, item_kind="DOCUMENT_CANDIDATE", readiness="ACTIONABLE", delivery_target="INTERNAL_REVIEW",
            deliverable_kind=request.deliverable_kind, input_refs=[f"work-plan-item:{request.work_plan_item_id}"],
            plan_version=1, profile_version=1, skill_id="test-skill", tool_id="test-tool")
        connection = MagicMock()
        connection.execute.return_value.fetchone.return_value = row
        with patch.object(delivery, "assert_case_work_plan_references_current") as plan_check, \
             patch.object(delivery, "_assert_authorized_source_manifest_is_current") as sources_check:
            result = delivery._read_staging_binding(connection, worker=worker, request=request)
            self.assertEqual(result["revision_request_id"], request.revision_request_id)
            plan_check.assert_called_once()
            sources_check.assert_called_once()
            sql, params = connection.execute.call_args.args
            self.assertIn("case_agent_document_content_generation_jobs inbox", sql)
            self.assertNotIn("case_agent_document_revision_inbox inbox", sql)
            self.assertIn("generation_review.candidate_hash = %s", sql)
            self.assertIn("inbox.lease_expires_at > clock_timestamp()", sql)
            self.assertIn("authority.revoked_at IS NULL", sql)
            self.assertEqual(params, (worker.actor_id, request.revision_request_id, worker.firm_id,
                request.content_generation_claim_version, request.candidate_hash, request.binding_hash))
            self.assertEqual(sql.count("%s"), len(params))
            # A content revision has a new lawyer generation-review gate.  It
            # must remain eligible when its already-verified predecessor task
            # was a read-only task with no original task approval gate.
            row["writes_managed_derivatives"] = False
            row["approval_gate"] = "NONE"
            delivery._read_staging_binding(connection, worker=worker, request=request)
            # A deterministic template reissue is also a forward-only
            # candidate.  With the exact source/plan/package binding intact,
            # it may re-render a package whose predecessor task was read-only.
            # The reissue cannot change content, approval, or submission state.
            row["inbox_state"] = "LEASED"
            delivery._read_staging_binding(connection, worker=worker,
                request=replace(request, generation_mode="DETERMINISTIC_TEMPLATE_REVISION", content_generation_claim_version=None))
            row["inbox_state"] = "RENDERING"
            row["writes_managed_derivatives"] = True
            row["approval_gate"] = "LAWYER_REVIEW"
            for key, value in (("inbox_state", "LEASED"), ("claimed_by", str(uuid4())), ("is_stale", True),
                               ("predecessor_revision_number", 2), ("plan_status", "DRAFT")):
                old = row[key]
                row[key] = value
                with self.subTest(key=key), self.assertRaises(CaseAgentDocumentPackageBlocked):
                    delivery._read_staging_binding(connection, worker=worker, request=request)
                row[key] = old
            row["inbox_state"] = "LEASED"
            delivery._read_staging_binding(connection, worker=worker,
                request=replace(request, generation_mode="DETERMINISTIC_TEMPLATE_REVISION", content_generation_claim_version=None))
            legacy_sql, legacy_params = connection.execute.call_args.args
            self.assertIn("case_agent_document_revision_inbox inbox", legacy_sql)
            self.assertNotIn("content_generation", legacy_sql)
            self.assertEqual(len(legacy_params), 3)

    def test_access_port_requires_different_system_workers(self):
        actor = Actor(str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER}))
        with self.assertRaisesRegex(ValueError, "must differ"):
            PostgresReviewableDocumentPackageAccessPort(
                dsn="postgresql://server-owned",
                verifier_actor=actor,
                execution_actor_id=actor.actor_id,
                object_store=_Objects(),
            )

    def test_access_port_exposes_the_unified_verifier_read_contract(self):
        actor = Actor(str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER}))
        port = PostgresReviewableDocumentPackageAccessPort(
            dsn="postgresql://server-owned",
            verifier_actor=actor,
            execution_actor_id=str(uuid4()),
            object_store=_Objects(),
        )
        selected = unittest.mock.Mock(
            artifact_id=str(uuid4()),
            artifact_kind="REVIEWABLE_DOCUMENT_EDITABLE",
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            content_sha256=digest(self.editable),
            byte_size=len(self.editable),
            content=self.editable,
        )
        package = unittest.mock.Mock(
            selected_artifact=selected,
            task_input_hash=self.task_hash,
            receipt_hash=digest("package receipt"),
        )
        artifact = ArtifactReceipt(
            artifact_id=selected.artifact_id,
            artifact_kind=selected.artifact_kind,
            content_hash=selected.content_sha256,
            byte_size=selected.byte_size,
            source_input_hash=self.task_hash,
            managed_derivative=True,
        )
        with patch.object(port, "read_package", return_value=package):
            managed = port.read_managed_artifact(
                firm_id=self.firm,
                matter_id=self.matter,
                run_id=self.run,
                artifact=artifact,
            )
        self.assertIsInstance(managed, ManagedArtifactRead)
        self.assertEqual(managed.content, self.editable)

    def test_first_release_verifier_has_all_three_document_package_formats(self):
        verifiers = first_release_review_candidate_verifiers()
        candidate = ManagedArtifactRead(
            artifact_id=str(uuid4()),
            artifact_kind="REVIEWABLE_DOCUMENT_CANDIDATE_JSON",
            content=self.candidate_bytes,
            source_input_hash=self.task_hash,
            object_receipt_hash=digest("candidate receipt"),
            media_type="application/json",
        )
        editable = ManagedArtifactRead(
            artifact_id=str(uuid4()),
            artifact_kind="REVIEWABLE_DOCUMENT_EDITABLE",
            content=self.editable,
            source_input_hash=self.task_hash,
            object_receipt_hash=digest("editable receipt"),
            media_type=self.request.editable_media_type,
        )
        preview = ManagedArtifactRead(
            artifact_id=str(uuid4()),
            artifact_kind="REVIEWABLE_DOCUMENT_PDF_PREVIEW",
            content=self.pdf,
            source_input_hash=self.task_hash,
            object_receipt_hash=digest("preview receipt"),
            media_type="application/pdf",
        )
        candidate_receipt = verifiers[candidate.artifact_kind].verify(candidate)
        self.assertEqual(
            candidate_receipt.declared_source_input_hash, self.task_hash
        )
        verifiers[editable.artifact_kind].verify(editable)
        verifiers[preview.artifact_kind].verify(preview)

    def test_document_aware_artifact_access_never_falls_back_for_documents(self):
        calls: list[str] = []

        class Access:
            def __init__(self, label): self.label = label
            def read_managed_artifact(self, **_kwargs):
                calls.append(self.label)
                return self.label

        router = DocumentAwareManagedArtifactAccessPort(
            base_access=Access("base"), document_access=Access("document")
        )
        document = ArtifactReceipt(
            artifact_id=str(uuid4()),
            artifact_kind="REVIEWABLE_DOCUMENT_EDITABLE",
            content_hash=digest("editable"),
            byte_size=1,
            source_input_hash=self.task_hash,
            managed_derivative=True,
        )
        self.assertEqual(
            router.read_managed_artifact(
                firm_id=self.firm,
                matter_id=self.matter,
                run_id=self.run,
                artifact=document,
            ),
            "document",
        )
        self.assertEqual(calls, ["document"])

    def test_document_delivery_preflight_accepts_exact_0039_contract(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        objects = _Objects()
        connection = _PreflightConnection()
        with patch(
            "case_kernel.case_agent_document_delivery_postgres._transaction",
            return_value=_Context(connection),
        ):
            preflight_case_agent_document_delivery_runtime_contract(
                dsn="postgresql://server-owned",
                worker_actor=worker,
                verifier_actor=verifier,
                object_store=objects,
            )
        self.assertEqual(objects.values, {})
        self.assertEqual(len(connection.executions), 3)

    def test_document_delivery_preflight_rejects_missing_table(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with patch(
            "case_kernel.case_agent_document_delivery_postgres._transaction",
            return_value=_Context(_PreflightConnection(omit_table=True)),
        ):
            with self.assertRaisesRegex(
                CaseAgentDocumentPackageBlocked, "0039 is incomplete"
            ):
                preflight_case_agent_document_delivery_runtime_contract(
                    dsn="postgresql://server-owned",
                    worker_actor=worker,
                    verifier_actor=verifier,
                    object_store=_Objects(),
                )

    def test_document_delivery_preflight_rejects_missing_key_column(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with patch(
            "case_kernel.case_agent_document_delivery_postgres._transaction",
            return_value=_Context(
                _PreflightConnection(omit_column="package_receipt_hash")
            ),
        ):
            with self.assertRaisesRegex(
                CaseAgentDocumentPackageBlocked, "0039 is incomplete"
            ):
                preflight_case_agent_document_delivery_runtime_contract(
                    dsn="postgresql://server-owned",
                    worker_actor=worker,
                    verifier_actor=verifier,
                    object_store=_Objects(),
                )

    def test_document_delivery_preflight_rejects_missing_guard(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        connection = _PreflightConnection(
            omit_trigger="case_agent_reviewable_document_packages_append_only"
        )
        with patch(
            "case_kernel.case_agent_document_delivery_postgres._transaction",
            return_value=_Context(connection),
        ):
            with self.assertRaisesRegex(
                CaseAgentDocumentPackageBlocked, "guards are incomplete"
            ):
                preflight_case_agent_document_delivery_runtime_contract(
                    dsn="postgresql://server-owned",
                    worker_actor=worker,
                    verifier_actor=verifier,
                    object_store=_Objects(),
                )

    def test_document_delivery_preflight_requires_force_rls(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with patch(
            "case_kernel.case_agent_document_delivery_postgres._transaction",
            return_value=_Context(_PreflightConnection(force_rls=False)),
        ):
            with self.assertRaisesRegex(
                CaseAgentDocumentPackageBlocked, "requires FORCE RLS"
            ):
                preflight_case_agent_document_delivery_runtime_contract(
                    dsn="postgresql://server-owned",
                    worker_actor=worker,
                    verifier_actor=verifier,
                    object_store=_Objects(),
                )

    def test_document_delivery_preflight_requires_distinct_system_workers(self):
        actor = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "must differ"
        ):
            preflight_case_agent_document_delivery_runtime_contract(
                dsn="postgresql://server-owned",
                worker_actor=actor,
                verifier_actor=actor,
                object_store=_Objects(),
            )

        human = Actor(str(uuid4()), self.firm, frozenset({Role.LEAD_LAWYER}))
        for worker, verifier in ((human, actor), (actor, human)):
            with self.subTest(worker=worker.roles, verifier=verifier.roles):
                with self.assertRaises(CaseAgentDocumentPackageBlocked):
                    preflight_case_agent_document_delivery_runtime_contract(
                        dsn="postgresql://server-owned",
                        worker_actor=worker,
                        verifier_actor=verifier,
                        object_store=_Objects(),
                    )

    def test_document_delivery_preflight_requires_private_store_contract(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with self.assertRaisesRegex(
            CaseAgentDocumentPackageBlocked, "preflight failed"
        ):
            preflight_case_agent_document_delivery_runtime_contract(
                dsn="postgresql://server-owned",
                worker_actor=worker,
                verifier_actor=verifier,
                object_store=object(),
            )

    def test_document_delivery_preflight_requires_buildable_template_registry(self):
        worker = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        verifier = Actor(
            str(uuid4()), self.firm, frozenset({Role.SYSTEM_WORKER})
        )
        with patch(
            "case_kernel.case_agent_document_delivery_postgres."
            "first_release_reviewable_document_templates",
            side_effect=ValueError("broken template"),
        ):
            with self.assertRaisesRegex(
                CaseAgentDocumentPackageBlocked, "preflight failed"
            ):
                preflight_case_agent_document_delivery_runtime_contract(
                    dsn="postgresql://server-owned",
                    worker_actor=worker,
                    verifier_actor=verifier,
                    object_store=_Objects(),
                )


if __name__ == "__main__":
    unittest.main()
