from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile
import unittest
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4

from case_kernel import approved_draft_worker as approved_draft_worker_module
from case_kernel.approved_draft_worker import (
    ApprovedDraft,
    ApprovedDraftBlocked,
    ApprovedSection,
    create_docx_draft,
    create_pdf_draft,
    create_xlsx_ledger,
)


class ApprovedDraftWorkerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.draft = ApprovedDraft(
            title="民事答辩状（草稿）",
            sections=(
                ApprovedSection(
                    "一、答辩意见",
                    ("对本金金额无异议。", "具体金额仍应以已确认交易台账及原始凭证为准。"),
                    ("事实-001", "规则-002"),
                ),
            ),
            approval_hash="a" * 64,
        )

    def test_docx_has_domestic_internal_legal_layout_and_source_binding(self) -> None:
        artifact = create_docx_draft(self.draft)
        self.assertEqual(
            artifact.media_type,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        document = Document(BytesIO(artifact.content))
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, places=1)
        self.assertAlmostEqual(section.page_height.mm, 297, places=1)
        self.assertAlmostEqual(section.top_margin.mm, 30, places=1)
        self.assertAlmostEqual(section.bottom_margin.mm, 25, places=1)
        self.assertAlmostEqual(section.left_margin.mm, 30, places=1)
        self.assertAlmostEqual(section.right_margin.mm, 25, places=1)
        self.assertAlmostEqual(section.header_distance.mm, 15, places=1)
        self.assertAlmostEqual(section.footer_distance.mm, 17.5, places=1)

        styles = document.styles
        self.assertEqual(approved_draft_worker_module._DOCX_TITLE_FONT, "黑体")
        self.assertEqual(approved_draft_worker_module._DOCX_BODY_FONT, "宋体")
        self.assertEqual(
            approved_draft_worker_module._DOCX_SECONDARY_HEADING_FONT,
            "楷体",
        )
        self.assertEqual(
            _east_asia_font(styles["Title"]),
            approved_draft_worker_module._DOCX_TITLE_FONT,
        )
        self.assertEqual(styles["Title"].font.size.pt, 18)
        self.assertEqual(str(styles["Title"].font.color.rgb), "000000")
        self.assertNotIn("w:pBdr", styles["Title"].element.xml)
        self.assertEqual(
            _east_asia_font(styles["Heading 1"]),
            approved_draft_worker_module._DOCX_HEADING_FONT,
        )
        self.assertEqual(styles["Heading 1"].font.size.pt, 14)
        self.assertEqual(str(styles["Heading 1"].font.color.rgb), "000000")
        self.assertEqual(
            _east_asia_font(styles["Heading 2"]),
            approved_draft_worker_module._DOCX_SECONDARY_HEADING_FONT,
        )
        self.assertEqual(styles["Heading 2"].font.size.pt, 12)
        self.assertFalse(styles["Heading 2"].font.bold)
        self.assertEqual(
            _east_asia_font(styles["Heading 4"]),
            approved_draft_worker_module._DOCX_BODY_FONT,
        )
        self.assertEqual(styles["Heading 4"].font.size.pt, 12)
        self.assertTrue(styles["Heading 4"].font.bold)
        self.assertFalse(styles["Heading 4"].font.italic)
        self.assertEqual(styles["Heading 4"].paragraph_format.line_spacing.pt, 22)
        self.assertTrue(styles["Heading 4"].paragraph_format.keep_with_next)
        self.assertIn('w:firstLineChars="0"', styles["Heading 4"].element.xml)
        self.assertEqual(
            _east_asia_font(styles["Normal"]),
            approved_draft_worker_module._DOCX_BODY_FONT,
        )
        for script_slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            self.assertEqual(
                _script_font(styles["Normal"], script_slot),
                approved_draft_worker_module._DOCX_BODY_FONT,
            )
            self.assertEqual(
                _script_font(styles["Title"], script_slot),
                approved_draft_worker_module._DOCX_TITLE_FONT,
            )
        self.assertEqual(styles["Normal"].font.size.pt, 12)
        self.assertEqual(styles["Normal"].paragraph_format.line_spacing.pt, 22)
        self.assertIn('w:firstLineChars="200"', styles["Normal"].element.xml)
        self.assertEqual(
            _east_asia_font(styles["来源注记"]),
            approved_draft_worker_module._DOCX_NOTE_FONT,
        )
        self.assertEqual(styles["来源注记"].font.size.pt, 10)

        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("律师复核候选｜待律师终审｜非正式文书", text)
        self.assertIn("版本摘要", text)
        self.assertNotIn("批准摘要", text)
        self.assertIn("生成时间以系统版本回执为准", text)
        self.assertIn("来源注记：事实-001；规则-002", text)
        self.assertEqual(document.paragraphs[-1].style.name, "来源注记")
        with ZipFile(BytesIO(artifact.content)) as archive:
            font_table = archive.read("word/fontTable.xml").decode("utf-8")
        for font_name in {
            approved_draft_worker_module._DOCX_TITLE_FONT,
            approved_draft_worker_module._DOCX_SECONDARY_HEADING_FONT,
            approved_draft_worker_module._DOCX_BODY_FONT,
        }:
            self.assertIn(f'w:name="{font_name}"', font_table)
        self.assertIn('<w:altName w:val="Heiti SC"/>', font_table)
        self.assertGreaterEqual(
            font_table.count('<w:altName w:val="Songti SC"/>'),
            2,
        )
        self.assertEqual(document.core_properties.identifier, self.draft.approval_hash)
        self.assertIn("律师复核候选｜待律师终审", section.header.paragraphs[0].text)
        self.assertIn("非正式文书｜律师复核候选", section.footer.paragraphs[0].text)
        self.assertNotIn("案件审阅意见", section.footer.paragraphs[0].text)
        self.assertIn(" PAGE ", section.footer.paragraphs[0]._p.xml)

    def test_docx_heading_prefixes_select_only_the_promised_visual_level(self) -> None:
        draft = ApprovedDraft(
            title="案件审阅意见",
            sections=(
                ApprovedSection("未编号概览", ("正文一。",), ("事实-001",)),
                ApprovedSection("（一）争议焦点", ("正文二。",), ("事实-002",)),
                ApprovedSection("1. 金额核对", ("正文三。",), ("事实-003",)),
                ApprovedSection("2、来源复核", ("正文四。",), ("事实-004",)),
                ApprovedSection("（1）凭证核对", ("正文五。",), ("事实-005",)),
                ApprovedSection("(2) 金额复算", ("正文六。",), ("事实-006",)),
            ),
            approval_hash="e" * 64,
        )
        document = Document(BytesIO(create_docx_draft(draft).content))
        styles_by_text = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs}
        self.assertEqual(styles_by_text["未编号概览"], "Heading 1")
        self.assertEqual(styles_by_text["（一）争议焦点"], "Heading 2")
        self.assertEqual(styles_by_text["1. 金额核对"], "Heading 3")
        self.assertEqual(styles_by_text["2、来源复核"], "Heading 3")
        self.assertEqual(styles_by_text["（1）凭证核对"], "Heading 4")
        self.assertEqual(styles_by_text["(2) 金额复算"], "Heading 4")

    def test_pdf_fallback_matches_a4_candidate_source_and_page_contract(self) -> None:
        artifact = create_pdf_draft(self.draft)
        self.assertEqual(artifact.media_type, "application/pdf")
        reader = PdfReader(BytesIO(artifact.content))
        self.assertEqual(len(reader.pages), 1)
        page = reader.pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), A4[0], places=1)
        self.assertAlmostEqual(float(page.mediabox.height), A4[1], places=1)
        extracted = "\n".join(current.extract_text() or "" for current in reader.pages)
        self.assertIn("律师复核候选｜待律师终审", extracted)
        self.assertIn("非正式文书｜律师复核候选｜第 1 页", extracted)
        self.assertNotIn("案件审阅意见", extracted)
        self.assertIn("来源注记：事实-001；规则-002", extracted)
        self.assertIn("版本摘要", extracted)
        self.assertNotIn("批准摘要", extracted)
        self.assertEqual(reader.metadata.title, self.draft.title)
        self.assertEqual(reader.metadata.subject, "国内律师内部复核候选（非正式文书）")
        fonts = [
            reference.get_object()
            for reference in page["/Resources"]["/Font"].values()
        ]
        self.assertGreaterEqual(len(fonts), 2)
        for font in fonts:
            descriptor = font["/FontDescriptor"].get_object()
            self.assertIn("/FontFile2", descriptor)
            self.assertIn("/ToUnicode", font)
            self.assertNotIn("Helvetica", str(font.get("/BaseFont")))
            self.assertNotIn("STSong-Light", str(font.get("/BaseFont")))

    def test_pdf_repeats_candidate_and_non_final_markers_on_every_page(self) -> None:
        long_draft = ApprovedDraft(
            title="案件审阅意见",
            sections=(
                ApprovedSection(
                    "一、事实核对",
                    tuple(f"第{index}项事实仅作律师复核候选，仍须回到原始材料核验。" for index in range(90)),
                    ("事实集合-001",),
                ),
            ),
            approval_hash="f" * 64,
        )
        reader = PdfReader(BytesIO(create_pdf_draft(long_draft).content))
        self.assertGreater(len(reader.pages), 1)
        for page_number, page in enumerate(reader.pages, start=1):
            extracted = page.extract_text() or ""
            self.assertIn("律师复核候选｜待律师终审", extracted)
            self.assertIn(f"非正式文书｜律师复核候选｜第 {page_number} 页", extracted)

    def test_pdf_wrap_keeps_legal_dates_amounts_and_ids_intact(self) -> None:
        approved_draft_worker_module._register_domestic_pdf_fonts()
        lines = approved_draft_worker_module._wrap_first_line(
            "金额￥310,638.59，期限2025-07-15，编号LAW-CIVIL-195，情景S-A-1。",
            150,
            12,
            0,
        )
        self.assertTrue(any("￥310,638.59" in line for line in lines))
        self.assertTrue(any("2025-07-15" in line for line in lines))
        self.assertTrue(any("LAW-CIVIL-195" in line for line in lines))
        self.assertTrue(any("S-A-1" in line for line in lines))
        self.assertFalse(any(line.startswith(("，", "。", "）")) for line in lines))

    def test_pdf_moves_a_short_paragraph_instead_of_orphaning_its_last_line(self) -> None:
        target = (
            "段落开始：该律师决定涉及事实真实性、证据对应关系和法律后果，"
            "应当作为一个完整复核单元呈现，不应把最后一行括号或结论单独留到下一页。"
            "段落结束。"
        )
        draft = ApprovedDraft(
            title="案件审阅意见",
            sections=(
                ApprovedSection(
                    "一、事实核对",
                    tuple(["填充段落。"] * 24 + [target]),
                    ("事实集合-001",),
                ),
            ),
            approval_hash="7" * 64,
        )
        pages = [
            page.extract_text() or ""
            for page in PdfReader(BytesIO(create_pdf_draft(draft).content)).pages
        ]
        start_page = next(index for index, text in enumerate(pages) if "段落开始" in text)
        end_page = next(index for index, text in enumerate(pages) if "段落结束" in text)
        self.assertEqual(start_page, end_page)

    def test_pdf_preserves_precomposed_latin_but_blocks_unshaped_or_unsafe_unicode(self) -> None:
        precomposed = ApprovedDraft(
            title="Café 案件审阅意见",
            sections=(
                ApprovedSection(
                    "一、原文核对",
                    ("Café 名称按原始预组字符保留。",),
                    ("事实-001",),
                ),
            ),
            approval_hash="9" * 64,
        )
        extracted = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(BytesIO(create_pdf_draft(precomposed).content)).pages
        )
        self.assertIn("Café", extracted)

        unsafe_samples = (
            ("Cafe\u0301", "0301"),  # detached combining mark
            ("AB\u200dCD", "200D"),  # zero-width joiner rendered as a box by UMing
            ("AB\u00adCD", "00AD"),  # soft hyphen rendered visibly by UMing
            ("材料\ufffc待解析", "FFFC"),  # unresolved embedded object marker
            ("姓名\ufffd", "FFFD"),  # decoding replacement marker
            ("姓名\U00020021", "20021"),  # ReportLab emits an invalid astral ToUnicode map
        )
        for unsafe_text, codepoint in unsafe_samples:
            with self.subTest(codepoint=codepoint):
                unsafe = ApprovedDraft(
                    title="案件审阅意见",
                    sections=(
                        ApprovedSection(
                            "一、原文核对",
                            (unsafe_text,),
                            ("事实-001",),
                        ),
                    ),
                    approval_hash="8" * 64,
                )
                with self.assertRaisesRegex(ApprovedDraftBlocked, rf"U\+{codepoint}"):
                    create_pdf_draft(unsafe)

    def test_pdf_font_selection_skips_empty_cmap_glyphs_before_rendering(self) -> None:
        codepoint = ord("℃")

        def fake_font(*, visible: bool) -> object:
            return SimpleNamespace(
                face=SimpleNamespace(
                    charWidths={codepoint: 1000},
                    charToGlyph={codepoint: 1},
                    glyphPos=[0, 0, 20 if visible else 0],
                )
            )

        fonts = {
            approved_draft_worker_module._PDF_HEADING_FONT: fake_font(visible=False),
            approved_draft_worker_module._PDF_TEXT_FONT: fake_font(visible=True),
            approved_draft_worker_module._PDF_SECONDARY_FONT: fake_font(visible=False),
        }
        with patch.object(
            approved_draft_worker_module.pdfmetrics,
            "getFont",
            side_effect=lambda name: fonts[name],
        ):
            self.assertEqual(
                approved_draft_worker_module._pdf_text_runs(
                    "℃", approved_draft_worker_module._PDF_HEADING_FONT
                ),
                ((approved_draft_worker_module._PDF_TEXT_FONT, "℃"),),
            )

        fonts[approved_draft_worker_module._PDF_TEXT_FONT] = fake_font(visible=False)
        with patch.object(
            approved_draft_worker_module.pdfmetrics,
            "getFont",
            side_effect=lambda name: fonts[name],
        ):
            with self.assertRaisesRegex(ApprovedDraftBlocked, r"U\+2103"):
                approved_draft_worker_module._pdf_text_runs(
                    "℃", approved_draft_worker_module._PDF_HEADING_FONT
                )

    def test_xlsx_ledger_has_chinese_review_and_print_layout_without_formulas(self) -> None:
        artifact = create_xlsx_ledger(
            approval_hash="b" * 64,
            sheet_name="付款台账",
            columns=("日期", "金额（原值）", "付款说明", "来源"),
            rows=(
                (
                    date(2024, 8, 15),
                    "1000.50",
                    '=HYPERLINK("https://example.invalid")',
                    "transaction:464f1ca7-4529-5851-b587-f3b2b79d354c",
                ),
                (date(2024, 8, 16), 2500, "  +SUM(1,1)", "事实-002"),
                (date(2024, 8, 17), "0.00", "-1+2", "事实-003"),
                (date(2024, 8, 18), "0.00", "@SUM(A1:A2)", "事实-004"),
            ),
        )
        workbook = load_workbook(BytesIO(artifact.content), data_only=False)
        sheet = workbook["付款台账"]
        self.assertEqual({str(item) for item in sheet.merged_cells.ranges}, {"A1:D1", "A2:D2"})
        self.assertEqual(sheet["A1"].value, "付款台账")
        self.assertIn("律师复核候选｜待律师终审｜非正式文书", sheet["A2"].value)
        self.assertIn("版本摘要", sheet["A2"].value)
        self.assertNotIn("批准摘要", sheet["A2"].value)
        self.assertEqual(tuple(cell.value for cell in sheet[3]), ("日期", "金额（原值）", "付款说明", "来源"))
        self.assertEqual(sheet.freeze_panes, "A4")
        self.assertEqual(sheet.auto_filter.ref, "A3:D7")
        self.assertEqual(sheet.print_title_rows, "$1:$3")
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(sheet.page_setup.paperSize, 9)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertIn("律师复核候选", sheet.oddHeader.center.text)
        self.assertIn("待律师终审｜非正式文书", sheet.oddFooter.left.text)
        self.assertIn("&P", sheet.oddFooter.center.text)

        self.assertEqual(sheet["A1"].font.name, "黑体")
        self.assertEqual(sheet["A1"].font.sz, 16)
        self.assertEqual(sheet["A3"].font.name, "黑体")
        self.assertEqual(sheet["A3"].font.sz, 11)
        self.assertEqual(sheet["A4"].font.name, "宋体")
        self.assertEqual(sheet["A4"].font.sz, 10.5)
        self.assertEqual(sheet["A4"].number_format, "yyyy-mm-dd")
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, 12)
        self.assertEqual(sheet["B4"].value, "1000.50")
        self.assertEqual(sheet["B4"].number_format, "@")
        self.assertEqual(sheet["B4"].alignment.horizontal, "right")
        self.assertEqual(sheet["B5"].number_format, "#,##0.00;[Red]-#,##0.00")
        self.assertEqual(sheet["C4"].value, '\'=HYPERLINK("https://example.invalid")')
        self.assertEqual(sheet["C5"].value, "'  +SUM(1,1)")
        self.assertEqual(sheet["C6"].value, "'-1+2")
        self.assertEqual(sheet["C7"].value, "'@SUM(A1:A2)")
        self.assertEqual(sheet.column_dimensions["D"].width, 18)
        self.assertGreaterEqual(sheet.row_dimensions[4].height, 45)
        self.assertEqual(sheet.row_dimensions[1].height, 30)
        self.assertEqual(workbook.calculation.calcMode, "manual")
        self.assertFalse(workbook.calculation.fullCalcOnLoad)
        self.assertFalse(workbook.calculation.forceFullCalc)
        self.assertFalse(workbook._external_links)
        self.assertFalse(
            any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row),
            "controlled ledgers must contain no executable cell formulas",
        )
        with ZipFile(BytesIO(artifact.content)) as archive:
            worksheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            core_properties = archive.read("docProps/core.xml").decode("utf-8")
        namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        self.assertEqual(worksheet.findall(".//x:f", namespace), [])
        self.assertEqual(core_properties.count("2000-01-01T00:00:00Z"), 2)

    def test_generated_artifact_bytes_and_hashes_are_deterministic(self) -> None:
        first_docx = create_docx_draft(self.draft)
        second_docx = create_docx_draft(self.draft)
        first_pdf = create_pdf_draft(self.draft)
        second_pdf = create_pdf_draft(self.draft)
        ledger_input = {
            "approval_hash": "c" * 64,
            "sheet_name": "付款台账",
            "columns": ("日期", "金额"),
            "rows": ((date(2024, 8, 15), "1000.50"),),
        }
        first_xlsx = create_xlsx_ledger(**ledger_input)
        second_xlsx = create_xlsx_ledger(**ledger_input)
        for first, second in (
            (first_docx, second_docx),
            (first_pdf, second_pdf),
            (first_xlsx, second_xlsx),
        ):
            self.assertEqual(first.content, second.content)
            self.assertEqual(first.content_sha256, second.content_sha256)

    def test_unapproved_or_unsafe_content_is_refused(self) -> None:
        with self.assertRaisesRegex(ApprovedDraftBlocked, "approval"):
            create_docx_draft(ApprovedDraft("草稿", self.draft.sections, "not-an-approval-hash"))
        with self.assertRaisesRegex(ApprovedDraftBlocked, "requires at least one"):
            create_xlsx_ledger(
                approval_hash="c" * 64,
                sheet_name="空台账",
                columns=("金额",),
                rows=(),
            )
        with self.assertRaisesRegex(ApprovedDraftBlocked, "cell value"):
            create_xlsx_ledger(
                approval_hash="d" * 64,
                sheet_name="付款台账",
                columns=("说明",),
                rows=(("危险\x00文本",),),
            )
        with self.assertRaisesRegex(ApprovedDraftBlocked, "printable wrapped-line"):
            create_xlsx_ledger(
                approval_hash="d" * 64,
                sheet_name="付款台账",
                columns=("来源",),
                rows=(("transaction:" + "a" * 500,),),
            )
        unsupported_pdf = ApprovedDraft(
            title="案件审阅意见",
            sections=(
                ApprovedSection(
                    "一、核对",
                    ("不得用缺字方框替代不受管的 emoji 💼。",),
                    ("事实-001",),
                ),
            ),
            approval_hash="e" * 64,
        )
        with self.assertRaisesRegex(ApprovedDraftBlocked, r"U\+1F4BC"):
            create_pdf_draft(unsupported_pdf)


def _east_asia_font(style: object) -> str | None:
    return style.element.rPr.rFonts.get(qn("w:eastAsia"))


def _script_font(style: object, script_slot: str) -> str | None:
    return style.element.rPr.rFonts.get(qn(script_slot))


if __name__ == "__main__":
    unittest.main()
