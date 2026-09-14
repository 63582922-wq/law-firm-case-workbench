from __future__ import annotations

from dataclasses import replace
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
import unittest
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from case_kernel.approved_draft_worker import create_xlsx_ledger
from case_kernel.case_agent_document_delivery import (
    AuthoritativeDocumentSource,
    CaseAgentDocumentDeliveryBlocked,
    DocumentSourceKind,
    DynamicDocumentTaskBinding,
    ReviewableDocumentFormat,
    build_document_draft_request,
    build_deterministic_case_review_memo_candidate,
    build_deterministic_defence_statement_candidate,
    build_deterministic_evidence_catalogue_candidate,
    build_deterministic_payment_ledger_candidate,
    build_deterministic_supplementary_evidence_checklist_candidate,
    canonical_document_candidate_bytes,
    first_release_reviewable_document_templates,
    parse_reviewable_document_candidate,
    visible_document_source_labels,
    _embeddable_clause,
)
from case_kernel.case_agent_lawyer_analysis import (
    build_lawyer_analysis_contract,
    compile_lawyer_decision_package_candidate,
    parse_lawyer_analysis_provider_response,
    prepare_lawyer_analysis_request,
)
from case_kernel.case_work_plan import (
    CaseWorkPlanItem,
    DeliveryTarget,
    ReviewGate,
    WorkPlanItemKind,
    WorkPlanReadiness,
)
from backend.tests.test_case_agent_lawyer_analysis import (
    ATTEMPT_ID as LAWYER_ATTEMPT_ID,
    HOST as LAWYER_HOST,
    _projection as lawyer_projection,
    _provider_response as lawyer_provider_response,
    _valid_core as valid_lawyer_core,
)


def digest(value: str) -> str:
    return sha256(value.encode()).hexdigest()


def lawyer_decision_package_payload() -> bytes:
    projection = lawyer_projection()
    contract, request = prepare_lawyer_analysis_request(
        projection=projection,
        task_id=projection.task_id,
        attempt_id=LAWYER_ATTEMPT_ID,
        endpoint_host=LAWYER_HOST,
    )
    parsed = parse_lawyer_analysis_provider_response(
        lawyer_provider_response(valid_lawyer_core(projection)),
        contract=contract,
    )
    return compile_lawyer_decision_package_candidate(
        projection=projection,
        contract=contract,
        parsed=parsed,
        external_request_id=request.external_request_id,
        request_hash=request.request_hash,
    )


class CaseAgentDocumentDeliveryTests(unittest.TestCase):
    def test_embeddable_clause_removes_stop_before_closing_quote(self):
        self.assertEqual(
            _embeddable_clause("律师如何处理“已确认事实。”"),
            "律师如何处理“已确认事实”",
        )

    def binding(self, deliverable: str = "DEFENCE_STATEMENT") -> DynamicDocumentTaskBinding:
        template = first_release_reviewable_document_templates().get(deliverable)
        is_defence = deliverable == "DEFENCE_STATEMENT"
        item = CaseWorkPlanItem(
            item_id=str(uuid4()),
            sequence=1,
            kind=WorkPlanItemKind.DOCUMENT_CANDIDATE,
            readiness=WorkPlanReadiness.ACTIONABLE,
            title="形成与当前诉请范围对应的文书候选",
            purpose="根据已确认来源形成供律师复核的文书候选。",
            rationale="当前程序、诉请范围和已确认法源支持形成候选。",
            prerequisites=(),
            trigger_refs=(),
            source_refs=(),
            risk_if_omitted="可能无法按期形成可复核文书。",
            confidence=0.91,
            review_gate=ReviewGate.LEAD_LAWYER_CONFIRMATION,
            delivery_target=(
                DeliveryTarget.INTERNAL_WORK_PRODUCT
                if is_defence
                else DeliveryTarget.COURT_SUBMISSION
            ),
            deliverable_kind=deliverable,
            required_for_delivery=True,
            is_primary_document=True,
        )
        source_kinds = set(template.required_source_kinds)
        sources = []
        for index, source_kind in enumerate(
            sorted(source_kinds, key=lambda value: value.value), start=1
        ):
            is_transaction = (
                source_kind is DocumentSourceKind.CONFIRMED_TRANSACTION
            )
            is_evidence_page = (
                source_kind is DocumentSourceKind.APPROVED_EVIDENCE_ITEM
            )
            source_text = f"仅用于测试的权威来源摘要 {index}"
            if source_kind is DocumentSourceKind.POSTURE_PROFILE:
                source_text = json.dumps(
                    {
                        "represented_party": "测试公司",
                        "represented_position": "DEFENDANT" if is_defence else "PLAINTIFF",
                        "procedure_stage": "FIRST_INSTANCE" if is_defence else "PRE_ACTION",
                        "case_type_code": "CIVIL.PRIVATE_LENDING" if is_defence else "SALE_CONTRACT_DISPUTE",
                        "authority_scope_code": "GENERAL_AUTHORITY" if is_defence else "LITIGATION_FULL",
                        "engagement_state": "ACTIVE",
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif is_evidence_page:
                source_text = json.dumps(
                    {
                        "page_number": 2,
                        "original_file_sha256": digest("evidence-page"),
                        "original_label": "付款凭证.pdf",
                        "disposition": "INCLUDE",
                        "reason": "与本案付款经过相关",
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif source_kind is DocumentSourceKind.WORK_PLAN_ITEM:
                source_text = json.dumps(
                    {
                        "title": "生成民事答辩状候选" if is_defence else "形成案件审阅意见候选",
                        "purpose": "将已确认来源编译为律师终审候选。" if is_defence else "识别案件风险并形成律师复核工作产品。",
                        "rationale": "当前已确认代理情境和来源已通过受控复核。" if is_defence else "当前已确认代理情境和核心事实。",
                        "risk_if_omitted": "律师可能无法按期完成逐项诉请复核。" if is_defence else "律师可能遗漏证据和程序风险。",
                        "delivery_target": "INTERNAL_WORK_PRODUCT" if is_defence else "INTERNAL_REVIEW",
                        "deliverable_kind": deliverable,
                        "required_for_delivery": True,
                        "is_primary_document": True,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif source_kind is DocumentSourceKind.CONFIRMED_FACT:
                source_text = "双方签订测试购销合同，付款期限已届满，现有材料包含一笔已确认付款。"
            elif source_kind is DocumentSourceKind.CONFIRMED_CLAIM:
                source_text = json.dumps(
                    {
                        "original_claim_text": "请求被告偿还借款本金人民币10,000.00元。",
                        "claimed_amount": "10000.00",
                        "currency": "CNY",
                        "position": "DISPUTE",
                        "partial_amount": None,
                        "partial_currency": None,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif source_kind is DocumentSourceKind.VERIFIED_LEGAL_SOURCE:
                source_text = json.dumps(
                    {
                        "publisher": "最高人民法院",
                        "authority_level": "司法解释",
                        "official_url": "https://www.court.gov.cn/example",
                        "provision_locator": "第六条",
                        "reviewed_text": "[合成] 已核验的官方法源摘录。",
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif source_kind is DocumentSourceKind.APPROVED_LEGAL_RULE:
                source_text = json.dumps(
                    {
                        "rule_id": "PRIVATE_LENDING_INTEREST",
                        "rule_version": "2026.01",
                        "issue_key": "PRIVATE_LENDING_INTEREST",
                        "effective_from": "2026-01-01",
                        "effective_to": None,
                        "trigger_event_kind": "DEFAULT",
                        "formula_kind": "FIXED_ANNUAL_RATE",
                        "base_annual_rate": "0.036500000000",
                        "rate_multiplier": None,
                        "derived_annual_rate": "0.036500000000",
                        "required_fact_keys": ["default_date"],
                        "transition_rule_versions": [],
                        "conflict_set": "PRIVATE_LENDING_RATE",
                        "priority": 1,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            elif (
                source_kind
                is DocumentSourceKind.VERIFIED_LAWYER_DECISION_PACKAGE
            ):
                source_text = lawyer_decision_package_payload().decode("utf-8")
            elif is_transaction:
                source_text = json.dumps(
                    {
                        "local_date": "2020-08-20",
                        "date_precision": "EXACT_DATE",
                        "amount": "1000.000000",
                        "currency": "CNY",
                        "direction": "OUTGOING",
                        "payer_label": "张三",
                        "payee_label": "李四",
                        "channel": "BANK",
                        "transaction_reference": "BANK-001",
                        "nature": "DISBURSEMENT",
                        "same_day_sequence": 1,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
            sources.append(
                AuthoritativeDocumentSource(
                    input_ref=(
                        f"transaction:{uuid4()}"
                        if is_transaction
                        else f"evidence-page:{uuid4()}"
                        if is_evidence_page
                        else f"claim:{uuid4()}"
                        if source_kind is DocumentSourceKind.CONFIRMED_CLAIM
                        else f"legal-source:{uuid4()}"
                        if source_kind is DocumentSourceKind.VERIFIED_LEGAL_SOURCE
                        else f"legal-rule:{uuid4()}"
                        if source_kind is DocumentSourceKind.APPROVED_LEGAL_RULE
                        else f"lawyer-decision-package:{uuid4()}"
                        if source_kind
                        is DocumentSourceKind.VERIFIED_LAWYER_DECISION_PACKAGE
                        else f"source:{index}"
                    ),
                    source_kind=source_kind,
                    source_version="v1",
                    source_hash=digest(source_kind.value),
                    label=f"来源 {index}",
                    text=source_text,
                )
            )
        return DynamicDocumentTaskBinding(
            firm_id=str(uuid4()),
            matter_id=str(uuid4()),
            run_id=str(uuid4()),
            graph_id=str(uuid4()),
            task_id=str(uuid4()),
            task_input_hash=digest("task"),
            case_snapshot_hash=digest("snapshot"),
            work_plan_id=str(uuid4()),
            work_plan_hash=digest("plan"),
            work_plan_status="ACTIVE",
            work_plan_item=item,
            posture_profile_id=str(uuid4()),
            posture_profile_hash=digest("posture"),
            template=template,
            sources=tuple(sources),
        )

    def response(self, binding: DynamicDocumentTaskBinding) -> bytes:
        common = {
            "binding": {
                "binding_hash": binding.binding_hash,
                "source_set_hash": binding.source_set_hash,
                "task_input_hash": binding.task_input_hash,
                "work_plan_item_id": binding.work_plan_item.item_id,
                "template_id": binding.template.template_id,
                "template_version": binding.template.template_version,
                "template_hash": binding.template.template_hash,
                "deliverable_kind": binding.template.deliverable_kind,
                "output_format": binding.template.output_format.value,
            },
            "title": binding.template.title_label,
            "review_status": "NEEDS_LAWYER_REVIEW",
            "formal_fact": False,
            "formal_legal_conclusion": False,
            "court_ready": False,
        }
        refs = [item.input_ref for item in binding.sources]
        if binding.template.output_format is ReviewableDocumentFormat.DOCX:
            value = {
                "schema_version": "case-agent-reviewable-docx-candidate-v1",
                **common,
                "sections": [
                    {
                        "heading": "待律师复核的回应",
                        "paragraphs": [
                            {
                                "text": "本段仅整理已确认来源，不代表最终法律意见。",
                                "source_refs": refs,
                            }
                        ],
                    }
                ],
            }
        else:
            value = {
                "schema_version": "case-agent-reviewable-xlsx-candidate-v1",
                **common,
                "columns": [
                    {"key": "date", "label": "日期", "value_type": "DATE"},
                    {"key": "amount", "label": "金额", "value_type": "DECIMAL"},
                ],
                "rows": [
                    {
                        "row_id": "row-1",
                        "cells": {"date": "2020-08-20", "amount": 1000.0},
                        "source_refs": refs,
                    }
                ],
            }
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()

    def test_first_release_templates_are_capabilities_not_role_workflows(self):
        templates = first_release_reviewable_document_templates()
        kinds = {item.deliverable_kind for item in templates.list_templates()}
        self.assertTrue({"COMPLAINT", "DEFENCE_STATEMENT", "APPEAL_PETITION"}.issubset(kinds))
        self.assertIn("SUPPLEMENTARY_EVIDENCE_CHECKLIST", kinds)
        self.assertEqual(templates.get("PAYMENT_LEDGER").output_format, ReviewableDocumentFormat.XLSX)
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "no installed"):
            templates.get("MODEL_INVENTED_DOCUMENT")

    def test_template_hash_binds_server_instructions_and_required_sources(self):
        template = first_release_reviewable_document_templates().get(
            "DEFENCE_STATEMENT"
        )
        self.assertEqual(len(template.template_hash), 64)
        changed = replace(
            template,
            drafting_instructions=template.drafting_instructions
            + ("新增但未升级版本的约束。",),
        )
        self.assertNotEqual(template.template_hash, changed.template_hash)
        changed_rendering = replace(
            template,
            rendering_instructions=template.rendering_instructions
            + ("新增但未升级版本的渲染约束。",),
        )
        self.assertNotEqual(template.template_hash, changed_rendering.template_hash)

    def test_domestic_template_release_is_versioned_and_hash_locked(self):
        expected_hashes = {
            "APPEAL_PETITION": "04ed90a0bcdcf0296948346cd8cc9239be3f2e7bb157b6669ad6ce0810d096c5",
            "CASE_REVIEW_MEMO": "a8873905b32f9051ae3a2a374232dc95dc32f63c7ce5123375cf8f6da2e0fb43",
            "COMPLAINT": "673449dd83575722649cf80b8a176fdabe8655ed78bcb4a0f3ecd917e3c5815a",
            "COUNTERCLAIM": "be5d311173ac97b18265c481a4cdfef050b9e29dc96b54eb1078cc652c2bf2a3",
            "DEFENCE_STATEMENT": "9f3d3846edabb134c4f54131e0b9e0c85c66ac41313983430cb1d4c0a7ec68e3",
            "EVIDENCE_CATALOGUE": "834985d3d458496bfbee42f984f1afbd932d369b9e8d1edf41bcf2a18e36f2b7",
            "INTEREST_CALCULATION_TABLE": "3ecc27620b3cc8cbde50a0ff833d30fcd515e47161bfeb9a15c87d90c4f1da7e",
            "LEGAL_RESEARCH_MEMO": "bb24b41fb58424d7a32ae9df24c6980f16042271ac9c980d696428cc07d044e9",
            "PAYMENT_LEDGER": "64485435e38de144189d27252f3d1459b8cc9e2bafb5976e8930a5e0571cf53a",
            "SUPPLEMENTARY_EVIDENCE_CHECKLIST": "1816a0528b7cfbe74444706b87b212d14f66bab23651556d62d2ed2b7977a466",
        }
        templates = first_release_reviewable_document_templates().list_templates()
        self.assertEqual(
            {item.deliverable_kind: item.template_hash for item in templates},
            expected_hashes,
        )
        self.assertEqual(
            {
                item.deliverable_kind: item.template_version
                for item in templates
                if item.output_format.value == "XLSX"
            },
            {
                "EVIDENCE_CATALOGUE": "1.2.2",
                "INTEREST_CALCULATION_TABLE": "1.2.2",
                "PAYMENT_LEDGER": "1.2.4",
            },
        )
        self.assertEqual(
            {
                item.template_version
                for item in templates
                if item.output_format.value != "XLSX"
                and item.deliverable_kind != "CASE_REVIEW_MEMO"
            },
            {"1.0.0", "1.1.2", "1.2.0"},
        )
        self.assertEqual(
            first_release_reviewable_document_templates()
            .get("CASE_REVIEW_MEMO")
            .template_version,
            "1.2.3",
        )

    def test_domestic_drafting_contract_covers_structure_numbers_law_and_status(self):
        registry = first_release_reviewable_document_templates()
        memo = registry.get("CASE_REVIEW_MEMO")
        memo_drafting = "\n".join(memo.drafting_instructions)
        memo_rendering = "\n".join(memo.rendering_instructions)
        for expected in (
            "当前可用边界—已确认事实",
            "第×条第×款第×项",
            "2024年3月5日",
            "人民币30,000.00元",
        ):
            self.assertIn(expected, memo_drafting)
        for expected in (
            "上页边距30毫米",
            "标题使用小二号（18磅）黑体",
            "二级标题使用小四号（12磅）楷体",
            "正文采用固定22磅行距",
            "律师复核候选｜待律师终审｜非正式文书",
            "不声称GB/T 9704—2012",
        ):
            self.assertIn(expected, memo_rendering)

        ledger = registry.get("PAYMENT_LEDGER")
        ledger_contract = "\n".join(
            ledger.drafting_instructions + ledger.rendering_instructions
        )
        for expected in (
            "列顺序由服务器固定",
            "标题16磅加粗",
            "冻结并自动筛选表头",
            "不启用公式、宏、外部链接",
        ):
            self.assertIn(expected, ledger_contract)

    def test_litigation_candidates_follow_supreme_court_public_section_order(self):
        registry = first_release_reviewable_document_templates()
        required_fragments = {
            "COMPLAINT": ("诉讼请求", "事实和理由", "证据和证据来源", "起诉人签名或盖章"),
            "DEFENCE_STATEMENT": ("案号和案由引言", "答辩意见", "证据和证据来源", "答辩人签名或盖章"),
            "COUNTERCLAIM": ("反诉请求", "事实和理由", "证据和证据来源", "反诉人签名或盖章"),
            "APPEAL_PETITION": ("原审法院、案号、案由", "上诉请求", "上诉理由", "上诉人签名或盖章"),
        }
        for deliverable_kind, fragments in required_fragments.items():
            instructions = "\n".join(
                registry.get(deliverable_kind).drafting_instructions
            )
            self.assertIn("最高人民法院公开的当事人参考民事诉讼文书样式", instructions)
            for fragment in fragments:
                self.assertIn(fragment, instructions)

    def test_server_request_is_hash_bound_and_contains_no_execution_controls(self):
        binding = self.binding()
        request = build_document_draft_request(binding)
        request.validate()
        value = json.loads(request.content)
        self.assertEqual(value["binding_hash"], binding.binding_hash)
        self.assertFalse(value["review_contract"]["court_ready"])
        self.assertEqual(value["rendering_contract"]["applied_by"], "SERVER_RENDERER")
        self.assertTrue(
            value["rendering_contract"][
                "model_must_not_simulate_layout_or_watermark_in_content"
            ]
        )
        self.assertEqual(
            value["rendering_contract"]["instructions"],
            list(binding.template.rendering_instructions),
        )
        encoded = request.content.decode()
        for forbidden in ('"file_path"', '"command"', '"tool_id"', '"provider"', '"endpoint"'):
            self.assertNotIn(forbidden, encoded)

    def test_docx_candidate_is_source_bound_review_only_and_renderable(self):
        binding = self.binding()
        candidate = parse_reviewable_document_candidate(self.response(binding), binding=binding)
        self.assertEqual(candidate.output_format, ReviewableDocumentFormat.DOCX)
        self.assertEqual(candidate.review_status, "NEEDS_LAWYER_REVIEW")
        draft = candidate.to_docx_input()
        self.assertEqual(draft.approval_hash, candidate.candidate_hash)
        self.assertEqual(set(draft.sections[0].source_refs), {item.input_ref for item in binding.sources})
        canonical = canonical_document_candidate_bytes(candidate)
        self.assertEqual(parse_reviewable_document_candidate(canonical, binding=binding), candidate)

    def test_lawyer_paragraph_preview_is_version_bound_and_not_approved(self):
        from dataclasses import replace
        from hashlib import sha256
        from case_kernel.case_agent_document_revisions import (
            LawyerParagraphChange, preview_lawyer_paragraph_changes,
            CaseAgentDocumentRevisionBlocked, CaseAgentDocumentRevisionConflict,
        )
        binding = self.binding()
        raw = self.response(binding)
        candidate = parse_reviewable_document_candidate(raw, binding=binding)
        original = candidate.sections[0].paragraphs[0]
        change = LawyerParagraphChange(0, 0, sha256(original.text.encode()).hexdigest(), "请进一步核对该项证据的证明范围。", "律师要求明确证据边界", original.source_refs)
        args = dict(binding=binding, current_candidate_bytes=raw, expected_candidate_hash=sha256(raw).hexdigest())
        preview = preview_lawyer_paragraph_changes(changes=(change,), **args)
        self.assertEqual(preview[0]["before"], original.text)
        self.assertEqual(preview[0]["after"], change.replacement_text)
        self.assertFalse(preview[0]["court_ready"])
        self.assertIn("修改前", preview[0]["diff"])
        self.assertEqual(candidate.sections[0].paragraphs[0], original)
        for invalid in (replace(change, section_index=-1), replace(change, source_refs=("foreign-source",)), replace(change, reason="")):
            with self.assertRaises(CaseAgentDocumentRevisionBlocked):
                preview_lawyer_paragraph_changes(changes=(invalid,), **args)
        with self.assertRaises(CaseAgentDocumentRevisionConflict):
            preview_lawyer_paragraph_changes(changes=(replace(change, expected_text_hash="0" * 64),), **args)
        with self.assertRaises(CaseAgentDocumentRevisionBlocked):
            preview_lawyer_paragraph_changes(changes=(change, change), **args)

    def test_lawyer_edit_prepares_a_distinct_unapproved_candidate(self):
        from hashlib import sha256
        from uuid import uuid4
        from case_kernel.models import Actor, Role
        from case_kernel.case_agent_document_revisions import LawyerParagraphChange, prepare_lawyer_document_revision
        binding = self.binding()
        raw = self.response(binding)
        original = parse_reviewable_document_candidate(raw, binding=binding)
        paragraph = original.sections[0].paragraphs[0]
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        change = LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(), "该项证据的证明范围仍需律师核对。", "明确证明边界", paragraph.source_refs)
        args = dict(actor=actor, binding=binding, changes=(change,), expected_candidate_hash=sha256(raw).hexdigest(), current_candidate_bytes=raw)
        prepared = prepare_lawyer_document_revision(**args)
        revised = parse_reviewable_document_candidate(prepared.candidate_content, binding=binding)
        self.assertEqual(revised.sections[0].paragraphs[0].text, change.replacement_text)
        self.assertNotEqual(revised.candidate_hash, original.candidate_hash)
        self.assertEqual(revised.review_status, "NEEDS_LAWYER_REVIEW")
        self.assertEqual(prepare_lawyer_document_revision(**args), prepared)
        self.assertEqual(sha256(prepared.review_manifest).hexdigest(), prepared.request_hash)
        self.assertFalse(json.loads(prepared.review_manifest)["court_ready"])
        self.assertEqual(json.loads(prepared.review_manifest)["matter_id"], binding.matter_id)
        self.assertEqual(json.loads(prepared.review_manifest)["run_id"], binding.run_id)
        self.assertEqual(json.loads(prepared.review_manifest)["task_id"], binding.task_id)
        self.assertFalse(json.loads(prepared.candidate_content)["formal_fact"])
        with self.assertRaises(PermissionError):
            prepare_lawyer_document_revision(**{**args, "actor": Actor(actor.actor_id, actor.firm_id, frozenset({Role.SYSTEM_WORKER}))})
        with self.assertRaisesRegex(PermissionError, "different firms"):
            prepare_lawyer_document_revision(**{**args, "actor": Actor(actor.actor_id, str(uuid4()), actor.roles)})

    def test_content_proposal_save_checks_current_package_and_never_enqueues(self):
        from hashlib import sha256
        from types import SimpleNamespace
        from uuid import uuid4
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        raw = self.response(binding)
        candidate = parse_reviewable_document_candidate(raw, binding=binding)
        paragraph = candidate.sections[0].paragraphs[0]
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        change = revisions.LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(), "该项证明范围待核验。", "律师修订", paragraph.source_refs)
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        store._read_state = Mock(return_value=SimpleNamespace(root_package_id=str(uuid4()), current_package_id=str(uuid4()), revision_number=1, version_status="CURRENT", run_status="READY_FOR_REVIEW"))
        connection = MagicMock()
        def execute(sql, params):
            if "FROM case_agent_document_content_proposals" in sql:
                return Mock(fetchone=Mock(return_value=None))
            if "SELECT binding_hash" in sql:
                return Mock(fetchone=Mock(return_value={"binding_hash": binding.binding_hash, "candidate_hash": candidate.candidate_hash}))
            return Mock()
        connection.execute.side_effect = execute
        context = MagicMock()
        context.__enter__.return_value = connection
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id, artifact_id=str(uuid4()), expected_revision_number=1, idempotency_key="content-proposal-test-0001", binding=binding, current_candidate_bytes=raw, expected_candidate_hash=sha256(raw).hexdigest(), changes=(change,))
        with patch.object(revisions, "_transaction", return_value=context):
            proposal_id = store.save_content_proposal(**args)
        self.assertTrue(proposal_id)
        sql = "\n".join(call.args[0] for call in connection.execute.call_args_list)
        self.assertIn("INSERT INTO case_agent_document_content_proposals", sql)
        self.assertNotIn("INSERT INTO case_agent_document_revision_requests", sql)
        self.assertNotIn("case_agent_document_revision_inbox", sql)
        saved = next(call.args[1] for call in connection.execute.call_args_list if "INSERT INTO case_agent_document_content_proposals" in call.args[0])
        prior = {"proposal_id": proposal_id, "request_hash": saved[8], "matter_id": binding.matter_id, "run_id": binding.run_id, "expected_revision_number": 1}
        connection.reset_mock()
        connection.execute.side_effect = lambda sql, params: Mock(fetchone=Mock(return_value=prior)) if "FROM case_agent_document_content_proposals" in sql else Mock()
        with patch.object(revisions, "_transaction", return_value=context):
            self.assertEqual(store.save_content_proposal(**args), proposal_id)
            prior["request_hash"] = "0" * 64
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                store.save_content_proposal(**args)
        self.assertFalse(any("INSERT INTO" in call.args[0] for call in connection.execute.call_args_list))

        connection.reset_mock()
        connection.execute.side_effect = execute
        store._read_state.return_value.revision_number = 2
        with patch.object(revisions, "_transaction", return_value=context):
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                store.save_content_proposal(**args)
        self.assertFalse(any("INSERT INTO" in call.args[0] for call in connection.execute.call_args_list))

    def test_content_revision_reconstruction_rejects_undeclared_edits_and_false_diff(self):
        import json
        from hashlib import sha256
        from uuid import uuid4
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        raw = self.response(binding)
        paragraph = parse_reviewable_document_candidate(raw, binding=binding).sections[0].paragraphs[0]
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        prepared = revisions.prepare_lawyer_document_revision(actor=actor, binding=binding,
            current_candidate_bytes=raw, expected_candidate_hash=sha256(raw).hexdigest(),
            changes=(revisions.LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(),
                     "该项证明范围待核验。", "律师修订", paragraph.source_refs),))
        args = dict(binding=binding, predecessor_content=raw, candidate_content=prepared.candidate_content,
                    review_manifest=prepared.review_manifest, request_hash=prepared.request_hash, requested_by=actor.actor_id)
        candidate = revisions.verify_lawyer_document_revision(**args)
        self.assertEqual(candidate.sections[0].paragraphs[0].text, "该项证明范围待核验。")
        from types import SimpleNamespace
        from unittest.mock import Mock, MagicMock, patch
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        state = SimpleNamespace(version_status="CURRENT", run_status="READY_FOR_REVIEW",
                                current_package_id=str(uuid4()), revision_number=1)
        store._read_state = Mock(return_value=state)
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        row = dict(candidate_content=prepared.candidate_content, review_manifest=prepared.review_manifest,
                   request_hash=prepared.request_hash, requested_by=actor.actor_id,
                   predecessor_candidate_hash=parse_reviewable_document_candidate(raw, binding=binding).candidate_hash)
        connection.execute.return_value.fetchone.return_value = row
        read_args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                         artifact_id=str(uuid4()), proposal_id=str(uuid4()), binding=binding, predecessor_content=raw)
        with patch.object(revisions, "_transaction", return_value=context) as transaction:
            self.assertEqual(store.read_verified_content_candidate(**read_args), candidate)
            transaction.assert_called_with("test-only", actor, read_only=True)
            row["predecessor_candidate_hash"] = "0" * 64
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_verified_content_candidate(**read_args)
            state.version_status = "UPDATE_REQUIRED"
            connection.reset_mock()
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_verified_content_candidate(**read_args)
            connection.execute.assert_not_called()
        with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
            revisions.verify_lawyer_document_revision(**{**args, "requested_by": str(uuid4())})
        with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
            revisions.verify_lawyer_document_revision(**{**args, "candidate_content": raw})
        for field, value in (("preview", []), ("binding_hash", "0" * 64), ("court_ready", True)):
            manifest = json.loads(prepared.review_manifest)
            manifest[field] = value
            altered = json.dumps(manifest, ensure_ascii=False).encode()
            # Even a recomputed storage digest cannot legitimize a false diff or changed scope.
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                revisions.verify_lawyer_document_revision(**{**args, "review_manifest": altered,
                                                              "request_hash": sha256(altered).hexdigest()})

    def test_content_generation_authorization_is_scoped_idempotent_and_not_submission(self):
        from types import SimpleNamespace
        from uuid import uuid4
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        state = SimpleNamespace(root_package_id=str(uuid4()), current_package_id=str(uuid4()), revision_number=1,
            template_id=binding.template.template_id, template_version=binding.template.template_version,
            template_hash=binding.template.template_hash, package_receipt_hash="b" * 64)
        store._read_state = Mock(return_value=state)
        store._verify_content_candidate = Mock(return_value=SimpleNamespace(candidate_hash="a" * 64))
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        authority, prior = {"ok": 1}, None
        def execute(sql, params):
            if "FROM users principal" in sql:
                return Mock(fetchone=Mock(return_value=authority))
            if "idempotency_key_hash = %s" in sql:
                return Mock(fetchone=Mock(return_value=prior))
            return Mock(fetchone=Mock(return_value=None))
        connection.execute.side_effect = execute
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()), proposal_id=str(uuid4()), expected_revision_number=1,
                    idempotency_key="document-review-test-0001", review_note="已核对修改及来源，仅生成复核稿。",
                    binding=binding, predecessor_content=self.response(binding))
        with patch.object(revisions, "_transaction", return_value=context):
            review_id = store.authorize_content_generation(**args)
            sql = "\n".join(call.args[0] for call in connection.execute.call_args_list)
            self.assertIn("INSERT INTO case_agent_document_content_generation_reviews", sql)
            self.assertIn("INSERT INTO case_agent_document_revision_requests", sql)
            registration = next(call.args[1] for call in connection.execute.call_args_list
                if "INSERT INTO case_agent_document_revision_requests" in call.args[0])
            self.assertEqual(registration[5:13], (state.root_package_id, state.current_package_id, 1,
                state.template_id, state.template_version, state.template_hash, state.package_receipt_hash, actor.actor_id))
            self.assertEqual(registration[-1], review_id)
            expected_hash = sha256("|".join(str(value) for value in (
                "case-agent-document-content-revision-request-v1", registration[0], actor.firm_id,
                binding.matter_id, binding.run_id, *registration[5:13], registration[1], review_id,
            )).encode()).hexdigest()
            self.assertEqual(registration[-2], expected_hash)
            self.assertNotIn("case_agent_document_content_generation_jobs", sql)
            self.assertNotIn("case_agent_document_revision_inbox", sql)
            self.assertNotIn("UPDATE", sql)
            store._verify_content_candidate.assert_called_once()
            prior = dict(review_id=review_id, proposal_id=args["proposal_id"], matter_id=binding.matter_id,
                         run_id=binding.run_id, root_package_id=state.root_package_id,
                         expected_revision_number=1, review_note=args["review_note"])
            connection.reset_mock()
            self.assertEqual(store.authorize_content_generation(**args), review_id)
            self.assertFalse(any("INSERT" in call.args[0] for call in connection.execute.call_args_list))
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                store.authorize_content_generation(**{**args, "review_note": "不同的决定理由"})
            prior = None
            with patch.object(store, "_register_content_revision_request", side_effect=RuntimeError("registration failed")):
                with self.assertRaisesRegex(RuntimeError, "registration failed"):
                    store.authorize_content_generation(**args)
            self.assertIs(context.__exit__.call_args.args[0], RuntimeError)
            state.revision_number = 2
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                store.authorize_content_generation(**args)
            authority = None
            with self.assertRaises(PermissionError):
                store.authorize_content_generation(**args)
            with self.assertRaises(PermissionError):
                store.authorize_content_generation(**{**args, "actor": Actor(actor.actor_id, actor.firm_id, frozenset({Role.ASSISTANT}))})

    def test_revision_request_hash_keeps_legacy_domain_and_binds_content_review(self):
        from case_kernel import case_agent_document_revisions as revisions
        values = dict(request_id=str(uuid4()), firm_id=str(uuid4()), matter_id=str(uuid4()), run_id=str(uuid4()),
            root_package_id=str(uuid4()), predecessor_package_id=str(uuid4()), expected_revision_number=1,
            target_template_id="test-template", target_template_version="1.0.0", target_template_hash="a" * 64,
            source_package_receipt_hash="b" * 64, requested_by=str(uuid4()), idempotency_key_hash="c" * 64)
        legacy = sha256("|".join(str(value) for value in (
            "case-agent-document-revision-request-v1", *values.values(),
        )).encode()).hexdigest()
        self.assertEqual(revisions._request_hash(**values), legacy)
        self.assertEqual(revisions._request_hash(**values, content_generation_review_id=None), legacy)
        review_id = str(uuid4())
        content = revisions._request_hash(**values, content_generation_review_id=review_id)
        self.assertNotEqual(content, legacy)
        self.assertNotEqual(content, revisions._request_hash(**values, content_generation_review_id=str(uuid4())))

    def test_content_revision_renderer_preserves_explicit_text_and_unknown_outcome(self):
        from zipfile import ZipFile
        from unittest.mock import Mock
        from case_kernel import case_agent_document_revisions as revisions
        from backend.tests.test_reviewable_draft_worker import FakeConverter
        binding = self.binding()
        payload = json.loads(self.response(binding))
        revised_text = "本段由律师修订，须核对原始凭证后再确认证明范围。"
        payload["sections"][0]["paragraphs"][0]["text"] = revised_text
        candidate = parse_reviewable_document_candidate(json.dumps(payload, ensure_ascii=False).encode(), binding=binding)
        converter = FakeConverter()
        result = revisions.render_content_revision_candidate(candidate=candidate, binding=binding, converter=converter)
        with ZipFile(BytesIO(result.editable_artifact.content)) as archive:
            self.assertIn(revised_text, archive.read("word/document.xml").decode())
        self.assertEqual(len(converter.calls), 1)
        self.assertEqual(result.review_pdf.source_sha256, result.editable_artifact.content_sha256)
        self.assertEqual(candidate.candidate_hash, result.approval_hash)
        unknown = Mock()
        unknown.convert_generated_document.side_effect = revisions.ReviewOfficeConversionUnknown("test timeout")
        with self.assertRaises(revisions.ReviewOfficeConversionUnknown):
            revisions.render_content_revision_candidate(candidate=candidate, binding=binding, converter=unknown)
        unknown.convert_generated_document.assert_called_once()
        rejected = Mock()
        with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
            revisions.render_content_revision_candidate(candidate=replace(candidate, candidate_hash="0" * 64), binding=binding, converter=rejected)
        rejected.convert_generated_document.assert_not_called()

    def test_authorized_content_read_checks_live_review_and_exact_candidate(self):
        from types import SimpleNamespace
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        candidate = parse_reviewable_document_candidate(self.response(binding), binding=binding)
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        state = SimpleNamespace(root_package_id=str(uuid4()), revision_number=1)
        store._read_state = Mock(return_value=state)
        store._verify_content_candidate = Mock(return_value=candidate)
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        row = dict(proposal_id=str(uuid4()), candidate_hash=candidate.candidate_hash)
        connection.execute.return_value.fetchone.return_value = row
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()), review_id=str(uuid4()), binding=binding, predecessor_content=self.response(binding))
        with patch.object(revisions, "_transaction", return_value=context):
            self.assertEqual(store.read_authorized_content_candidate(**args), candidate)
            sql, params = connection.execute.call_args.args
            self.assertIn("assignment.revoked_at IS NULL", sql)
            self.assertIn("review.purpose = 'GENERATE_REVIEW_COPY'", sql)
            self.assertEqual(params[-2:], (1, binding.binding_hash))
            row["candidate_hash"] = "0" * 64
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_authorized_content_candidate(**args)
            connection.execute.return_value.fetchone.return_value = None
            store._verify_content_candidate.reset_mock()
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_authorized_content_candidate(**args)
            store._verify_content_candidate.assert_not_called()

    def test_claimed_content_uses_worker_identity_and_replays_exact_authorized_edit(self):
        from unittest.mock import MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions

        binding = self.binding()
        raw = self.response(binding)
        predecessor = parse_reviewable_document_candidate(raw, binding=binding)
        paragraph = predecessor.sections[0].paragraphs[0]
        lawyer = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        worker = Actor(str(uuid4()), binding.firm_id, frozenset({Role.SYSTEM_WORKER}))
        prepared = revisions.prepare_lawyer_document_revision(actor=lawyer, binding=binding,
            current_candidate_bytes=raw, expected_candidate_hash=sha256(raw).hexdigest(),
            changes=(revisions.LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(),
                "该项证明范围待律师核验。", "保留证据边界", paragraph.source_refs),))
        expected = parse_reviewable_document_candidate(prepared.candidate_content, binding=binding)
        row = dict(candidate_content=prepared.candidate_content, review_manifest=prepared.review_manifest,
            request_hash=prepared.request_hash, requested_by=lawyer.actor_id,
            predecessor_candidate_hash=predecessor.candidate_hash, authorized_candidate_hash=expected.candidate_hash)
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        connection.execute.return_value.fetchone.return_value = row
        jobs = revisions.PostgresContentGenerationJobStore(dsn="test-only", worker_actor=worker)
        cursor = connection.execute.return_value
        connection.execute.side_effect = lambda sql, *params: (
            MagicMock(fetchone=MagicMock(return_value={"recovery_results_available": True})) if "to_regclass" in sql else cursor)
        claim = revisions.ContentGenerationClaim(str(uuid4()), binding.matter_id, binding.run_id, 2)
        with patch.object(revisions, "_transaction", return_value=context) as transaction:
            actual = jobs.read_claimed_candidate(claim, binding=binding, predecessor_content=raw)
            self.assertEqual(actual, expected)
            transaction.assert_called_with("test-only", worker, read_only=True)
            sql, params = connection.execute.call_args.args
            self.assertEqual(params, (claim.review_id, worker.firm_id, binding.matter_id, binding.run_id,
                worker.actor_id, 2, binding.binding_hash, binding.task_id))
            for predicate in ("job.lease_expires_at > now()", "job.state = 'LEASED'",
                    "receipt.outcome = 'PASSED'", "run.snapshot_matter_version = matter.version",
                    "assignment.revoked_at IS NULL", "predecessor.revision_number = review.expected_revision_number"):
                self.assertIn(predicate, sql)
            self.assertFalse(any("UPDATE " in call.args[0] or "INSERT " in call.args[0]
                for call in connection.execute.call_args_list))
            for key in ("predecessor_candidate_hash", "authorized_candidate_hash", "request_hash"):
                original = row[key]
                row[key] = "0" * 64
                with self.subTest(key=key), self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                    jobs.read_claimed_candidate(claim, binding=binding, predecessor_content=raw)
                row[key] = original
            connection.execute.return_value.fetchone.return_value = None
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                jobs.read_claimed_candidate(claim, binding=binding, predecessor_content=raw)
            transaction.reset_mock()
            for invalid in (replace(claim, run_id=str(uuid4())), replace(claim, matter_id=str(uuid4())),
                            replace(claim, claim_version=True), replace(claim, claim_version=0)):
                with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                    jobs.read_claimed_candidate(invalid, binding=binding, predecessor_content=raw)
            transaction.assert_not_called()

    def test_content_consumer_preserves_edit_and_stops_unknown_after_render(self):
        from types import SimpleNamespace
        from unittest.mock import Mock, patch
        from io import BytesIO
        from docx import Document
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        from backend.tests.test_reviewable_draft_worker import FakeConverter
        binding = self.binding()
        raw = self.response(binding)
        paragraph = parse_reviewable_document_candidate(raw, binding=binding).sections[0].paragraphs[0]
        lawyer = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        prepared = revisions.prepare_lawyer_document_revision(actor=lawyer, binding=binding,
            current_candidate_bytes=raw, expected_candidate_hash=sha256(raw).hexdigest(),
            changes=(revisions.LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(),
                "该项证明范围待律师核验。", "补充审阅边界", paragraph.source_refs),))
        candidate = parse_reviewable_document_candidate(prepared.candidate_content, binding=binding)
        claim = revisions.ContentGenerationClaim(str(uuid4()), binding.matter_id, binding.run_id, 2)
        revision = revisions.DocumentRevisionClaim(str(uuid4()), "a" * 64, str(uuid4()), str(uuid4()), 1,
            lawyer.actor_id, binding.run_id, binding.matter_id, str(uuid4()), binding.task_id, str(uuid4()), binding.task_input_hash, 2)
        loaded = revisions.ClaimedContentRevision(revision, str(uuid4()))
        predecessor = SimpleNamespace(package_id=revision.predecessor_package_id, revision_number=1,
            candidate=SimpleNamespace(content=raw))
        package = SimpleNamespace(package_id=str(uuid4()), receipt_hash="b" * 64, generation_mode="LAWYER_CONTENT_REVISION",
            revision_request_id=revision.request_id, supersedes_package_id=revision.predecessor_package_id,
            root_package_id=revision.root_package_id, revision_number=2, content_generation_claim_version=2,
            candidate_hash=candidate.candidate_hash)
        staged = SimpleNamespace(package_id=package.package_id, receipt_hash=package.receipt_hash,
            candidate_artifact=SimpleNamespace(artifact_id=str(uuid4())))
        events = []
        worker = revisions.PostgresContentRevisionWorker.__new__(revisions.PostgresContentRevisionWorker)
        worker._worker = Actor(str(uuid4()), binding.firm_id, frozenset({Role.SYSTEM_WORKER}))
        worker._jobs = SimpleNamespace(claim=Mock(side_effect=[claim, None]), load_revision=Mock(return_value=loaded),
            read_claimed_candidate=Mock(return_value=candidate), begin_render=Mock(side_effect=lambda value: events.append("begin")),
            mark_unknown=Mock())
        worker._binding = SimpleNamespace(resolve_document_revision=Mock(return_value=binding))
        worker._converter = FakeConverter()
        def stage(request):
            events.append("stage")
            self.assertEqual(request.content_generation_claim_version, 2)
            self.assertEqual(request.generation_mode, "LAWYER_CONTENT_REVISION")
            self.assertEqual(request.candidate_bytes, canonical_document_candidate_bytes(candidate))
            text = "\n".join(p.text for p in Document(BytesIO(request.editable_bytes)).paragraphs)
            self.assertIn("该项证明范围待律师核验。", text)
            return staged
        worker._packages = SimpleNamespace(stage_package=Mock(side_effect=stage))
        def read(**kwargs):
            if kwargs["artifact_id"] == loaded.predecessor_candidate_artifact_id:
                events.append("read-original")
                return predecessor
            events.append("verify-new")
            return package
        worker._access = SimpleNamespace(read_package=Mock(side_effect=read))
        worker._receipts = SimpleNamespace(record=Mock(side_effect=lambda **kwargs: events.append(kwargs["outcome"])))
        self.assertTrue(worker.run_cycle())
        self.assertFalse(worker.run_cycle())
        self.assertEqual(events, ["read-original", "begin", "stage", "verify-new", "PASSED"])
        worker._jobs.mark_unknown.assert_not_called()
        worker._binding.resolve_document_revision.assert_called_once_with(request_id=revision.request_id,
            predecessor_package_id=revision.predecessor_package_id, expected_revision_number=1, content_claim_version=2)
        for failure, expected in ((revisions.ReviewOfficeConversionUnknown("unknown"), "UNKNOWN"),
                                  (RuntimeError("private provider error"), "UNKNOWN")):
            worker._jobs.claim.side_effect = [claim, None]
            worker._packages.stage_package.reset_mock()
            worker._receipts.record.reset_mock()
            with patch.object(revisions, "render_content_revision_candidate", side_effect=failure) as renderer:
                self.assertTrue(worker.run_cycle())
                self.assertFalse(worker.run_cycle())
                renderer.assert_called_once()
            worker._packages.stage_package.assert_not_called()
            self.assertEqual(worker._receipts.record.call_args.kwargs["outcome"], expected)
            self.assertIsNone(worker._receipts.record.call_args.kwargs["successor_package_id"])
        worker._jobs.claim.side_effect = [claim, None]
        worker._receipts.record.reset_mock()
        worker._receipts.record.side_effect = [RuntimeError("receipt acknowledgement lost"), None]
        worker._packages.stage_package.reset_mock()
        self.assertTrue(worker.run_cycle())
        self.assertFalse(worker.run_cycle())
        worker._packages.stage_package.assert_called_once()
        self.assertEqual(worker._receipts.record.call_count, 2)
        self.assertEqual(*worker._receipts.record.call_args_list)
        worker._jobs.mark_unknown.assert_not_called()
        worker._jobs.claim.side_effect = [claim, None]
        worker._receipts.record.reset_mock()
        worker._receipts.record.side_effect = RuntimeError("receipt acknowledgement lost")
        worker._packages.stage_package.reset_mock()
        self.assertTrue(worker.run_cycle())
        self.assertFalse(worker.run_cycle())
        worker._packages.stage_package.assert_called_once()
        self.assertEqual(worker._receipts.record.call_count, 2)
        self.assertEqual(*worker._receipts.record.call_args_list)
        worker._jobs.mark_unknown.assert_called_once_with(claim)
        worker._jobs.claim.side_effect = [claim, None]
        worker._receipts.record.reset_mock()
        worker._jobs.mark_unknown.reset_mock()
        worker._receipts.record.side_effect = revisions.CaseAgentDocumentRevisionConflict("immutable UNKNOWN")
        self.assertTrue(worker.run_cycle())
        worker._receipts.record.assert_called_once()
        worker._jobs.mark_unknown.assert_called_once_with(claim)
        worker._jobs.claim.side_effect = [claim, None]
        worker._receipts.record.side_effect = None
        worker._receipts.record.reset_mock()
        package.candidate_hash = "0" * 64
        self.assertTrue(worker.run_cycle())
        self.assertFalse(worker.run_cycle())
        self.assertEqual(worker._receipts.record.call_args.kwargs["outcome"], "FAILED")
        self.assertIsNone(worker._receipts.record.call_args.kwargs["successor_package_id"])
        package.candidate_hash = candidate.candidate_hash
        worker._jobs.claim.side_effect = [claim, None]
        worker._packages.stage_package.side_effect = RuntimeError("registration result unknown")
        worker._receipts.record.reset_mock()
        self.assertTrue(worker.run_cycle())
        self.assertFalse(worker.run_cycle())
        self.assertEqual(worker._receipts.record.call_args.kwargs["outcome"], "UNKNOWN")

    def test_content_generation_job_claim_fences_render_and_never_reclaims_rendering(self):
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        actor = Actor(str(uuid4()), str(uuid4()), frozenset({Role.SYSTEM_WORKER}))
        jobs = revisions.PostgresContentGenerationJobStore(dsn="test-only", worker_actor=actor)
        row = dict(review_id=str(uuid4()), matter_id=str(uuid4()), run_id=str(uuid4()), claim_version=2)
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        connection.execute.return_value.fetchone.return_value = row
        with patch.object(revisions, "_transaction", return_value=context):
            claim = jobs.claim()
            self.assertEqual(claim.claim_version, 2)
            cleanup, claiming = [call.args[0] for call in connection.execute.call_args_list]
            self.assertIn("WHEN state = 'RENDERING' THEN 'UNKNOWN'", cleanup)
            self.assertIn("claim_version < 3", claiming)
            self.assertNotIn("RENDERING", claiming)
            self.assertIn("FOR UPDATE SKIP LOCKED", claiming)
            jobs.begin_render(claim)
            sql, params = connection.execute.call_args.args
            self.assertIn("claimed_by = %s AND claim_version = %s AND state = %s", sql)
            self.assertEqual(params[-4:], (actor.actor_id, 2, "LEASED", True))
            jobs.mark_unknown(claim)
            self.assertEqual(connection.execute.call_args.args[1][-2:], ("RENDERING", False))
            connection.execute.return_value.fetchone.return_value = None
            with self.assertRaises(revisions.CaseAgentDocumentRevisionConflict):
                jobs.begin_render(claim)
            self.assertIsNone(jobs.claim())
        with self.assertRaises(PermissionError):
            revisions.PostgresContentGenerationJobStore(dsn="test-only", worker_actor=Actor(actor.actor_id, actor.firm_id, frozenset({Role.LEAD_LAWYER})))

    def test_content_proposal_resolution_is_read_only_and_actor_scoped(self):
        from hashlib import sha256
        from types import SimpleNamespace
        from uuid import uuid4
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        root, proposal_id = str(uuid4()), str(uuid4())
        store._read_state = Mock(return_value=SimpleNamespace(root_package_id=root))
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        connection.execute.return_value.fetchone.return_value = {"proposal_id": proposal_id}
        key = "document-edit-test-0001"
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()), idempotency_key=key)
        with patch.object(revisions, "_transaction", return_value=context) as transaction:
            self.assertEqual(store.resolve_content_proposal(**args), proposal_id)
            transaction.assert_called_with("test-only", actor, read_only=True)
            sql, params = connection.execute.call_args.args
            self.assertIn("proposal.requested_by = %s", sql)
            self.assertNotIn("INSERT", sql)
            self.assertNotIn("candidate_content", sql)
            self.assertEqual(params, (actor.firm_id, binding.matter_id, binding.run_id, actor.actor_id, sha256(key.encode()).hexdigest(), root))
            connection.execute.return_value.fetchone.return_value = None
            self.assertIsNone(store.resolve_content_proposal(**args))
            connection.reset_mock()
            store._read_state.side_effect = revisions.CaseAgentDocumentRevisionBlocked("denied")
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.resolve_content_proposal(**args)
            connection.execute.assert_not_called()

    def test_content_proposal_list_is_bounded_scoped_and_does_not_read_body(self):
        from datetime import datetime, timezone
        from types import SimpleNamespace
        from uuid import uuid4
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        root = str(uuid4())
        store._read_state = Mock(return_value=SimpleNamespace(root_package_id=root))
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        rows = [dict(proposal_id=str(uuid4()), expected_revision_number=1,
                     created_at=datetime.now(timezone.utc)) for _ in range(21)]
        connection.execute.return_value.fetchall.return_value = rows
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()))
        with patch.object(revisions, "_transaction", return_value=context) as transaction:
            result = store.list_content_proposals(**args)
            transaction.assert_called_with("test-only", actor, read_only=True)
            self.assertEqual(len(result["items"]), 20)
            self.assertEqual(result["next_after"], rows[19]["proposal_id"])
            sql, params = connection.execute.call_args.args
            self.assertIn("LIMIT 21", sql)
            self.assertIn("ORDER BY created_at DESC, proposal_id DESC", sql)
            self.assertIn("FROM scoped_proposals WHERE proposal_id = %s::uuid", sql)
            self.assertNotIn("review_manifest", sql)
            self.assertNotIn("candidate_content", sql)
            self.assertEqual(params, (actor.firm_id, binding.matter_id, binding.run_id, root, None, None))
            connection.execute.return_value.fetchall.return_value = []
            self.assertEqual(store.list_content_proposals(**args, after=rows[19]["proposal_id"]), {"items": [], "next_after": None})
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.list_content_proposals(**args, after="not-a-cursor")
            connection.reset_mock()
            store._read_state.side_effect = revisions.CaseAgentDocumentRevisionBlocked("denied")
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.list_content_proposals(**args)
            connection.execute.assert_not_called()

    def test_saved_content_proposal_read_is_scoped_read_only_and_marks_old_version(self):
        from datetime import datetime, timezone
        from hashlib import sha256
        from types import SimpleNamespace
        from uuid import uuid4
        from unittest.mock import Mock, MagicMock, patch
        from case_kernel.models import Actor, Role
        from case_kernel import case_agent_document_revisions as revisions
        binding = self.binding()
        raw = self.response(binding)
        paragraph = parse_reviewable_document_candidate(raw, binding=binding).sections[0].paragraphs[0]
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        prepared = revisions.prepare_lawyer_document_revision(
            actor=actor, binding=binding, current_candidate_bytes=raw,
            expected_candidate_hash=sha256(raw).hexdigest(),
            changes=(revisions.LawyerParagraphChange(0, 0, sha256(paragraph.text.encode()).hexdigest(),
                     "该项证明范围待核验。", "律师修订", paragraph.source_refs),),
        )
        store = revisions.PostgresDocumentRevisionCommandStore.__new__(revisions.PostgresDocumentRevisionCommandStore)
        store._dsn = "test-only"
        package_id, proposal_id = str(uuid4()), str(uuid4())
        state = SimpleNamespace(root_package_id=package_id, current_package_id=package_id,
                                revision_number=1, version_status="CURRENT")
        store._read_state = Mock(return_value=state)
        row = dict(proposal_id=proposal_id, predecessor_package_id=package_id,
                   expected_revision_number=1, requested_by=actor.actor_id,
                   created_at=datetime.now(timezone.utc), request_hash=prepared.request_hash,
                   review_manifest=prepared.review_manifest)
        connection, context = MagicMock(), MagicMock()
        context.__enter__.return_value = connection
        connection.execute.return_value.fetchone.return_value = row
        args = dict(actor=actor, matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()), proposal_id=proposal_id)
        cursor = connection.execute.return_value
        connection.execute.side_effect = lambda sql, *params: (
            MagicMock(fetchone=MagicMock(return_value={"recovery_results_available": True})) if "to_regclass" in sql else cursor)
        with patch.object(revisions, "_transaction", return_value=context) as transaction:
            result = store.read_content_proposal(**args)
            transaction.assert_called_with("test-only", actor, read_only=True)
            self.assertTrue(result["based_on_current_version"])
            self.assertFalse(result["court_ready"])
            self.assertEqual(result["changes"][0]["before"], paragraph.text)
            self.assertNotIn("binding_hash", result)
            self.assertNotIn("candidate_content", result)
            sql, params = next(call.args for call in connection.execute.call_args_list
                               if "COALESCE(predecessor.root_package_id, predecessor.package_id)" in call.args[0])
            self.assertIn("COALESCE(predecessor.root_package_id, predecessor.package_id)", sql)
            self.assertEqual(params, (proposal_id, actor.firm_id, binding.matter_id, binding.run_id, package_id))
            state.current_package_id, state.revision_number = str(uuid4()), 2
            self.assertFalse(store.read_content_proposal(**args)["based_on_current_version"])
            row["request_hash"] = "0" * 64
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_content_proposal(**args)
            row["request_hash"] = prepared.request_hash
            connection.execute.return_value.fetchone.return_value = None
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_content_proposal(**args)
            connection.reset_mock()
            store._read_state.side_effect = revisions.CaseAgentDocumentRevisionBlocked("access denied")
            with self.assertRaises(revisions.CaseAgentDocumentRevisionBlocked):
                store.read_content_proposal(**args)
            connection.execute.assert_not_called()

    def test_content_proposal_read_service_authorizes_before_reading_content(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        from uuid import uuid4
        from case_kernel.models import Actor, Role
        from case_api.web_case_agent_documents import PostgresWebCaseAgentDocumentReviewService, WebCaseAgentDocumentReviewBlocked
        binding = self.binding()
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        service = PostgresWebCaseAgentDocumentReviewService.__new__(PostgresWebCaseAgentDocumentReviewService)
        service._read_revision_state = Mock(return_value=(actor, object()))
        result = {"proposal_id": str(uuid4()), "court_ready": False}
        read = Mock(return_value=result)
        service._revisions = SimpleNamespace(read_content_proposal=read)
        args = dict(identity=object(), matter_id=binding.matter_id, run_id=binding.run_id,
                    artifact_id=str(uuid4()), proposal_id=result["proposal_id"])
        self.assertEqual(service.read_content_proposal(**args), result)
        read.assert_called_once_with(actor=actor, **{key: value for key, value in args.items() if key != "identity"})
        read.reset_mock()
        service._read_revision_state.side_effect = PermissionError("denied")
        with self.assertRaises(PermissionError):
            service.read_content_proposal(**args)
        read.assert_not_called()
        service._read_revision_state.side_effect = None
        service._revisions = object()
        with self.assertRaises(WebCaseAgentDocumentReviewBlocked):
            service.read_content_proposal(**args)

    def test_content_proposal_service_uses_only_verified_server_bytes(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        from uuid import uuid4
        from case_kernel.models import Actor, Role
        from case_api.web_case_agent_documents import PostgresWebCaseAgentDocumentReviewService, WebCaseAgentDocumentReviewBlocked
        binding = self.binding()
        raw = self.response(binding)
        actor = Actor(str(uuid4()), binding.firm_id, frozenset({Role.LEAD_LAWYER}))
        service = PostgresWebCaseAgentDocumentReviewService.__new__(PostgresWebCaseAgentDocumentReviewService)
        state = SimpleNamespace(version_status="CURRENT", revision_number=1)
        package = SimpleNamespace(output_format=ReviewableDocumentFormat.DOCX, task_id=binding.task_id, binding_hash=binding.binding_hash, candidate=SimpleNamespace(content=raw))
        service._read_revision_state = Mock(return_value=(actor, state))
        service._read_verified_package = Mock(return_value=package)
        service._content_binding_resolver = Mock(return_value=binding)
        service._revisions = SimpleNamespace(save_content_proposal=Mock(return_value=str(uuid4())))
        args = dict(identity=object(), matter_id=binding.matter_id, run_id=binding.run_id, artifact_id=str(uuid4()), expected_revision_number=1, idempotency_key="content-proposal-test-0001", changes=())
        service.save_content_proposal(**args)
        self.assertEqual(service._revisions.save_content_proposal.call_args.kwargs["current_candidate_bytes"], raw)
        service._revisions.save_content_proposal.reset_mock()
        service._content_binding_resolver = None
        with self.assertRaisesRegex(WebCaseAgentDocumentReviewBlocked, "尚未配置"):
            service.save_content_proposal(**args)
        service._revisions.save_content_proposal.assert_not_called()

        service._content_binding_resolver = Mock(return_value=binding)
        authorize = Mock(return_value=str(uuid4()))
        service._revisions.authorize_content_generation = authorize
        review_args = {key: value for key, value in args.items() if key != "changes"}
        review_args.update(proposal_id=str(uuid4()), review_note="已核对本次修改及来源。")
        service.authorize_content_generation(**review_args)
        self.assertEqual(authorize.call_args.kwargs["predecessor_content"], raw)
        authorize.reset_mock()
        service._read_revision_state.return_value = (Actor(actor.actor_id, actor.firm_id, frozenset({Role.ASSISTANT})), state)
        with self.assertRaises(PermissionError):
            service.authorize_content_generation(**review_args)
        authorize.assert_not_called()

    def test_xlsx_candidate_never_emits_formula_and_appends_source_column(self):
        binding = self.binding("PAYMENT_LEDGER")
        candidate = parse_reviewable_document_candidate(self.response(binding), binding=binding)
        sheet, columns, rows = candidate.to_xlsx_input()
        self.assertLessEqual(len(sheet), 31)
        self.assertEqual(columns[-1], "来源")
        self.assertTrue(rows[0][-1])
        value = json.loads(self.response(binding))
        value["rows"][0]["cells"]["date"] = "=HYPERLINK(\"https://example.invalid\")"
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "formula-like"):
            parse_reviewable_document_candidate(json.dumps(value).encode(), binding=binding)

    def test_payment_ledger_is_one_exact_server_row_per_confirmed_transaction(self):
        binding = self.binding("PAYMENT_LEDGER")
        candidate = build_deterministic_payment_ledger_candidate(binding)
        transaction = next(
            source
            for source in binding.sources
            if source.source_kind is DocumentSourceKind.CONFIRMED_TRANSACTION
        )
        self.assertEqual(len(candidate.rows), 1)
        self.assertEqual(candidate.rows[0].row_id, transaction.input_ref)
        self.assertEqual(candidate.rows[0].source_refs, (transaction.input_ref,))
        self.assertEqual(
            dict(
                zip(
                    (column.key for column in candidate.columns),
                    candidate.rows[0].cells,
                    strict=True,
                )
            ),
            json.loads(transaction.text),
        )

    def test_evidence_catalogue_only_projects_approved_pages_and_leaves_review_judgment_open(self):
        binding = self.binding("EVIDENCE_CATALOGUE")
        candidate = build_deterministic_evidence_catalogue_candidate(binding)
        source = next(
            item
            for item in binding.sources
            if item.source_kind is DocumentSourceKind.APPROVED_EVIDENCE_ITEM
        )
        self.assertEqual(candidate.deliverable_kind, "EVIDENCE_CATALOGUE")
        self.assertEqual(candidate.rows[0].row_id, source.input_ref)
        self.assertEqual(candidate.rows[0].source_refs, (source.input_ref,))
        columns = {column.key: index for index, column in enumerate(candidate.columns)}
        self.assertEqual(candidate.rows[0].cells[columns["proof_purpose"]], "待律师核对")
        self.assertEqual(candidate.rows[0].cells[columns["review_status"]], "已纳入材料范围；三性待律师核对")
        self.assertEqual(
            parse_reviewable_document_candidate(
                canonical_document_candidate_bytes(candidate), binding=binding
            ),
            candidate,
        )

    def test_case_review_memo_is_deterministic_substantive_and_source_bound(self):
        binding = self.binding("CASE_REVIEW_MEMO")
        candidate = build_deterministic_case_review_memo_candidate(binding)
        self.assertEqual(candidate.deliverable_kind, "CASE_REVIEW_MEMO")
        self.assertEqual(len(candidate.sections), 10)
        content = "\n".join(
            paragraph.text
            for section in candidate.sections
            for paragraph in section.paragraphs
        )
        for expected in (
            "争点与证据风险",
            "对方可能主张与反制",
            "策略路径与取舍",
            "没有已确认诉请",
            "没有已批准计算结果",
            "没有已核验法源",
            "必须由律师决定",
            "可选处理",
            "解除暂缓条件线索",
        ):
            self.assertIn(expected, "\n".join((*(s.heading for s in candidate.sections), content)))
        self.assertNotIn("。。", content)
        self.assertNotIn("。”。", content)
        self.assertIn("行动1（现在处理）：关联风险", content)
        self.assertIn("责任：律师团队", content)
        self.assertIn("律师决定1（关联风险", content)
        self.assertNotIn("Agent候选建议：Agent ", content)
        self.assertNotIn("真实调用和计费", content)
        self.assertNotIn("模型运行", "\n".join(section.heading for section in candidate.sections))
        self.assertNotIn("行动1（现在处理）：围绕“围绕", content)
        self.assertNotIn("律师决定1：律师如何处理“围绕", content)
        allowed_refs = {source.input_ref for source in binding.sources}
        self.assertTrue(
            all(
                paragraph.source_refs
                and set(paragraph.source_refs).issubset(allowed_refs)
                for section in candidate.sections
                for paragraph in section.paragraphs
            )
        )
        self.assertEqual(
            parse_reviewable_document_candidate(
                canonical_document_candidate_bytes(candidate), binding=binding
            ),
            candidate,
        )

    def test_supplementary_evidence_checklist_is_source_bound_and_does_not_claim_collection(self):
        binding = self.binding("SUPPLEMENTARY_EVIDENCE_CHECKLIST")
        candidate = build_deterministic_supplementary_evidence_checklist_candidate(binding)
        self.assertEqual(candidate.deliverable_kind, "SUPPLEMENTARY_EVIDENCE_CHECKLIST")
        self.assertEqual(
            tuple(section.heading for section in candidate.sections),
            (
                "一、清单使用说明",
                "二、优先补齐材料",
                "三、当事人待答问题",
                "四、取得与核对行动",
                "五、提交前核查",
            ),
        )
        content = "\n".join(
            paragraph.text
            for section in candidate.sections
            for paragraph in section.paragraphs
        )
        self.assertIn("补证事项1", content)
        self.assertIn("当事人待答问题1", content)
        self.assertIn("取得与核对行动1", content)
        self.assertIn("不等于已取得证据", content)
        self.assertNotIn("模型运行", content)
        allowed_refs = {source.input_ref for source in binding.sources}
        self.assertTrue(
            all(
                paragraph.source_refs
                and set(paragraph.source_refs).issubset(allowed_refs)
                for section in candidate.sections
                for paragraph in section.paragraphs
            )
        )
        self.assertEqual(
            parse_reviewable_document_candidate(
                canonical_document_candidate_bytes(candidate), binding=binding
            ),
            candidate,
        )

    def test_defence_statement_is_deterministic_source_bound_and_review_only(self):
        binding = self.binding("DEFENCE_STATEMENT")
        candidate = build_deterministic_defence_statement_candidate(binding)
        self.assertEqual(candidate.deliverable_kind, "DEFENCE_STATEMENT")
        self.assertEqual(candidate.review_status, "NEEDS_LAWYER_REVIEW")
        self.assertEqual(len(candidate.sections), 7)
        content = "\n".join(
            paragraph.text
            for section in candidate.sections
            for paragraph in section.paragraphs
        )
        for expected in (
            "对原告诉请的逐项回应",
            "处理口径为提出争议",
            "已核验法源",
            "已批准规则",
            "待律师编目",
            "系统不会自动填写、批准、锁定或发送",
        ):
            self.assertIn(expected, "\n".join((*(s.heading for s in candidate.sections), content)))
        self.assertNotIn("法院：", content)
        self.assertNotIn("案号：", content)
        allowed_refs = {source.input_ref for source in binding.sources}
        self.assertTrue(
            all(
                paragraph.source_refs
                and set(paragraph.source_refs).issubset(allowed_refs)
                for section in candidate.sections
                for paragraph in section.paragraphs
            )
        )
        self.assertEqual(
            parse_reviewable_document_candidate(
                canonical_document_candidate_bytes(candidate), binding=binding
            ),
            candidate,
        )

    def test_defence_statement_fails_closed_for_wrong_posture_or_unbound_claim_response(self):
        binding = self.binding("DEFENCE_STATEMENT")
        posture = next(
            source
            for source in binding.sources
            if source.source_kind is DocumentSourceKind.POSTURE_PROFILE
        )
        wrong_posture = replace(
            posture,
            text=json.dumps(
                {**json.loads(posture.text), "represented_position": "PLAINTIFF"},
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
        )
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "defendant posture"):
            build_deterministic_defence_statement_candidate(
                replace(
                    binding,
                    sources=tuple(
                        wrong_posture if source is posture else source
                        for source in binding.sources
                    ),
                )
            )
        claim = next(
            source
            for source in binding.sources
            if source.source_kind is DocumentSourceKind.CONFIRMED_CLAIM
        )
        missing_response = replace(
            claim,
            text=json.dumps(
                {**json.loads(claim.text), "position": "UNDECIDED"},
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
        )
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "response position"):
            build_deterministic_defence_statement_candidate(
                replace(
                    binding,
                    sources=tuple(
                        missing_response if source is claim else source
                        for source in binding.sources
                    ),
                )
            )

    def test_real_payment_ledger_chain_projects_typed_lawyer_usable_cells(self):
        binding = self.binding("PAYMENT_LEDGER")
        candidate = build_deterministic_payment_ledger_candidate(binding)
        sheet_name, columns, rows = candidate.to_xlsx_input(
            visible_document_source_labels(binding)
        )
        artifact = create_xlsx_ledger(
            approval_hash=candidate.candidate_hash,
            sheet_name=sheet_name,
            columns=columns,
            rows=rows,
        )
        workbook = load_workbook(BytesIO(artifact.content), data_only=False)
        sheet = workbook[sheet_name]
        values = dict(zip(columns, (cell.value for cell in sheet[4]), strict=True))
        self.assertEqual(values["日期"], date(2020, 8, 20))
        self.assertEqual(values["日期精度"], "精确到日")
        self.assertEqual(values["金额"], 1000)
        self.assertEqual(values["币种"], "人民币（CNY）")
        self.assertEqual(values["收付方向"], "付款")
        self.assertEqual(values["渠道"], "银行转账")
        self.assertEqual(values["交易性质"], "款项交付")
        self.assertNotIn("transaction:", values["来源"])
        self.assertIn("已确认交易", values["来源"])
        self.assertEqual(
            sheet.cell(row=4, column=columns.index("日期") + 1).data_type,
            "d",
        )
        self.assertEqual(
            sheet.cell(row=4, column=columns.index("金额") + 1).data_type,
            "n",
        )
        self.assertEqual(
            sheet.cell(row=4, column=columns.index("金额") + 1).number_format,
            "#,##0.00;[Red]-#,##0.00",
        )
        self.assertEqual(
            sheet.cell(row=4, column=columns.index("金额") + 1).alignment.horizontal,
            "right",
        )
        expected_widths = {
            "日期": 15,
            "日期精度": 12,
            "收付方向": 11,
            "付款方": 14,
            "收款方": 14,
            "交易性质": 15,
            "来源": 18,
        }
        for label, expected_width in expected_widths.items():
            with self.subTest(column=label):
                letter = get_column_letter(columns.index(label) + 1)
                self.assertEqual(
                    sheet.column_dimensions[letter].width, expected_width
                )
        self.assertGreaterEqual(sheet.row_dimensions[4].height, 45)
        sequence_cell = sheet.cell(
            row=4, column=columns.index("同日顺序") + 1
        )
        source_cell = sheet.cell(row=4, column=columns.index("来源") + 1)
        self.assertEqual(sequence_cell.alignment.horizontal, "center")
        self.assertEqual(source_cell.alignment.horizontal, "left")
        self.assertEqual(source_cell.border.left.style, "thin")
        self.assertFalse(
            any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)
        )

    def test_payment_ledger_rejects_noncanonical_or_missing_transaction_source(self):
        binding = self.binding("PAYMENT_LEDGER")
        transaction = next(
            source
            for source in binding.sources
            if source.source_kind is DocumentSourceKind.CONFIRMED_TRANSACTION
        )
        tampered = replace(
            transaction,
            text=transaction.text.replace('"1000.000000"', '"-999"'),
        )
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "amount"):
            build_deterministic_payment_ledger_candidate(
                replace(
                    binding,
                    sources=tuple(
                        tampered if source is transaction else source
                        for source in binding.sources
                    ),
                )
            )
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "lacks sources"):
            build_deterministic_payment_ledger_candidate(
                replace(
                    binding,
                    sources=tuple(
                        source
                        for source in binding.sources
                        if source.source_kind
                        is not DocumentSourceKind.CONFIRMED_TRANSACTION
                    ),
                )
            )

    def test_candidate_rejects_unknown_source_and_formal_claims(self):
        binding = self.binding()
        value = json.loads(self.response(binding))
        value["sections"][0]["paragraphs"][0]["source_refs"] = ["source:999"]
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "outside"):
            parse_reviewable_document_candidate(json.dumps(value).encode(), binding=binding)
        value = json.loads(self.response(binding))
        value["court_ready"] = True
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "formal or court-ready"):
            parse_reviewable_document_candidate(json.dumps(value).encode(), binding=binding)
        value = json.loads(self.response(binding))
        value["title"] = "已正式出具的法律意见"
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "authorized template"):
            parse_reviewable_document_candidate(json.dumps(value).encode(), binding=binding)

    def test_candidate_rejects_duplicate_keys_and_task_drift(self):
        binding = self.binding()
        raw = self.response(binding).decode()
        malformed = raw[:-1] + ',"title":"重复标题"}'
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "strict JSON"):
            parse_reviewable_document_candidate(malformed, binding=binding)
        stale = replace(binding, work_plan_status="SUPERSEDED")
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "active work plan"):
            parse_reviewable_document_candidate(self.response(binding), binding=stale)

    def test_binding_requires_sources_declared_by_template(self):
        binding = self.binding("INTEREST_CALCULATION_TABLE")
        incomplete = replace(
            binding,
            sources=tuple(
                item
                for item in binding.sources
                if item.source_kind is not DocumentSourceKind.APPROVED_CALCULATION
            ),
        )
        with self.assertRaisesRegex(CaseAgentDocumentDeliveryBlocked, "lacks sources"):
            build_document_draft_request(incomplete)


if __name__ == "__main__":
    unittest.main()
