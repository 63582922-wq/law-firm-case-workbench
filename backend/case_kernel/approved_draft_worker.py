"""Create Word, Excel and PDF drafts from lawyer-approved structured content.

This is an output worker, not a general file editor.  It does not open or
overwrite source documents: the caller supplies an immutable approved content
snapshot and stores the resulting bytes in the encrypted managed artifact
area.  Formula execution, macros, external links and arbitrary file paths are
outside this worker's contract.
"""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
import re
from typing import Iterable

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


class ApprovedDraftBlocked(ValueError):
    """A proposed document is not an approved bounded draft input."""


@dataclass(frozen=True)
class ApprovedSection:
    heading: str
    paragraphs: tuple[str, ...]
    source_refs: tuple[str, ...]


@dataclass(frozen=True)
class ApprovedDraft:
    title: str
    sections: tuple[ApprovedSection, ...]
    approval_hash: str


@dataclass(frozen=True)
class DraftArtifact:
    media_type: str
    content: bytes
    content_sha256: str


_MAX_SECTIONS = 200
_MAX_PARAGRAPHS = 20_000
_MAX_TEXT_CHARS = 1_000_000
_MAX_SHEETS = 100
_MAX_ROWS = 100_000
_MAX_COLUMNS = 200
_MAX_LEDGER_TEXT_CHARS = 10_000_000
_TEXT_FONT = "STSong-Light"
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")


def create_docx_draft(draft: ApprovedDraft) -> DraftArtifact:
    _validate_draft(draft)
    document = Document()
    document.core_properties.title = draft.title
    document.add_heading(draft.title, level=0)
    for section in draft.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        document.add_paragraph("来源：" + "；".join(section.source_refs)).italic = True
    destination = BytesIO()
    document.save(destination)
    return _artifact("application/vnd.openxmlformats-officedocument.wordprocessingml.document", destination.getvalue())


def create_xlsx_ledger(
    *,
    approval_hash: str,
    sheet_name: str,
    columns: tuple[str, ...],
    rows: Iterable[tuple[str | int | float | None, ...]],
) -> DraftArtifact:
    _validate_approval_hash(approval_hash)
    if not sheet_name.strip() or len(sheet_name) > 31 or any(char in sheet_name for char in "[]:*?/\\"):
        raise ApprovedDraftBlocked("ledger sheet name is invalid")
    if not 1 <= len(columns) <= _MAX_COLUMNS or any(not value.strip() or len(value) > 160 for value in columns):
        raise ApprovedDraftBlocked("ledger columns are invalid")
    validated_rows: list[tuple[str | int | float | None, ...]] = []
    row_count = 0
    text_characters = sum(len(column) for column in columns)
    for row in rows:
        if len(row) != len(columns):
            raise ApprovedDraftBlocked("ledger row width does not match columns")
        values = tuple(_safe_cell_value(value) for value in row)
        text_characters += sum(len(value) for value in values if isinstance(value, str))
        if text_characters > _MAX_LEDGER_TEXT_CHARS:
            raise ApprovedDraftBlocked("ledger text exceeds the supported boundary")
        validated_rows.append(values)
        row_count += 1
        if row_count > _MAX_ROWS:
            raise ApprovedDraftBlocked("ledger row limit exceeded")
    if row_count == 0:
        raise ApprovedDraftBlocked("ledger requires at least one approved row")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=sheet_name.strip())
    sheet.append(list(columns))
    for row in validated_rows:
        sheet.append(row)
    destination = BytesIO()
    workbook.save(destination)
    return _artifact("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", destination.getvalue())


def create_pdf_draft(draft: ApprovedDraft) -> DraftArtifact:
    _validate_draft(draft)
    pdfmetrics.registerFont(UnicodeCIDFont(_TEXT_FONT))
    destination = BytesIO()
    document = canvas.Canvas(destination, pagesize=A4, pageCompression=1, invariant=1)
    document.setTitle(draft.title)
    width, height = A4
    left = 54
    y = height - 56
    document.setFont(_TEXT_FONT, 16)
    document.drawString(left, y, draft.title)
    y -= 30
    for section in draft.sections:
        y = _ensure_y(document, y, minimum=60)
        document.setFont(_TEXT_FONT, 12)
        document.drawString(left, y, section.heading)
        y -= 20
        document.setFont(_TEXT_FONT, 10)
        for paragraph in section.paragraphs:
            for line in _wrap(paragraph, width - left * 2, 10):
                y = _ensure_y(document, y, minimum=54)
                document.drawString(left, y, line)
                y -= 15
        document.setFont(_TEXT_FONT, 8)
        for line in _wrap("来源：" + "；".join(section.source_refs), width - left * 2, 8):
            y = _ensure_y(document, y, minimum=54)
            document.drawString(left, y, line)
            y -= 12
        y -= 8
    document.save()
    return _artifact("application/pdf", destination.getvalue())


def _validate_draft(draft: ApprovedDraft) -> None:
    _validate_approval_hash(draft.approval_hash)
    if not draft.title.strip() or len(draft.title) > 240:
        raise ApprovedDraftBlocked("draft title is invalid")
    if not 1 <= len(draft.sections) <= _MAX_SECTIONS:
        raise ApprovedDraftBlocked("draft section count is invalid")
    paragraphs = 0
    characters = len(draft.title)
    for section in draft.sections:
        if not section.heading.strip() or len(section.heading) > 240 or not section.source_refs:
            raise ApprovedDraftBlocked("each section requires a heading and approved source references")
        for source_ref in section.source_refs:
            if not source_ref.strip() or len(source_ref) > 240:
                raise ApprovedDraftBlocked("draft source reference is invalid")
        for paragraph in section.paragraphs:
            if not paragraph.strip():
                raise ApprovedDraftBlocked("draft paragraph is blank")
            paragraphs += 1
            characters += len(paragraph)
    if paragraphs < 1 or paragraphs > _MAX_PARAGRAPHS or characters > _MAX_TEXT_CHARS:
        raise ApprovedDraftBlocked("draft content exceeds the supported boundary")


def _validate_approval_hash(value: str) -> None:
    if not _SHA256_RE.fullmatch(value):
        raise ApprovedDraftBlocked("draft requires a lawyer approval SHA-256")


def _safe_cell_value(value: str | int | float | None) -> str | int | float | None:
    if value is None or isinstance(value, (int, float)):
        return value
    if not isinstance(value, str) or len(value) > 10_000:
        raise ApprovedDraftBlocked("ledger cell value is invalid")
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _ensure_y(document: canvas.Canvas, y: float, *, minimum: float) -> float:
    if y >= minimum:
        return y
    document.showPage()
    return A4[1] - 56


def _wrap(value: str, width: float, size: int) -> tuple[str, ...]:
    lines: list[str] = []
    current = ""
    for char in value.replace("\t", "    "):
        candidate = current + char
        if current and pdfmetrics.stringWidth(candidate, _TEXT_FONT, size) > width:
            lines.append(current)
            current = char
        else:
            current = candidate
    return tuple(lines + [current]) if current else tuple(lines or [""])


def _artifact(media_type: str, content: bytes) -> DraftArtifact:
    if not content:
        raise ApprovedDraftBlocked("draft generator returned an empty artifact")
    return DraftArtifact(media_type, content, sha256(content).hexdigest())
